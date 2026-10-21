"""Демо-процесс «светофор договоров» (templates/processes/svetofor-dogovorov) на файле кейса.

Запуск: python -m unittest discover -s tools/tests -p "test_*.py"
Эталон Практики 22.10: дата расчёта 21.09.2026, жёлтый порог 0-30 дней, 25 строк на трёх листах.
"""

import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent.parent
DEMO = ROOT / "templates" / "processes" / "svetofor-dogovorov"
CASE_FILE = ROOT / "input" / "case-azs" / "data" / "reestr_dogovorov.xlsx"
SLUG = "svetofor-dogovorov"


def load_demo():
    spec = importlib.util.spec_from_file_location("svetofor_demo", DEMO / "run.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class DemoNumbersTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.mod = load_demo()
        cls.res = cls.mod.build_result(CASE_FILE)
        cls.col = {name: i for i, name in enumerate(cls.res["columns"])}

    def numbers(self, status):
        i = self.col["№ договора"]
        return sorted(r[i] for r, s in zip(self.res["rows"], self.res["row_status"]) if s == status)

    def row(self, number):
        i = self.col["№ договора"]
        found = [(r, s, n) for r, s, n in zip(self.res["rows"], self.res["row_status"], self.res["row_note"]) if r[i] == number]
        self.assertTrue(found, number)
        return found

    def test_date_is_legend_constant(self):
        self.assertEqual(self.mod.CALC_DATE, date(2026, 9, 21))
        self.assertEqual(self.mod.YELLOW_DAYS, 30)
        self.assertIn("21.09.2026", self.res["summary"])

    def test_counts_match_reference(self):
        self.assertEqual(len(self.res["rows"]), 25)
        self.assertEqual(
            Counter(self.res["row_status"]),
            Counter({"green": 16, "red": 3, "yellow": 3, "grey": 2, "info": 1}),
        )
        self.assertEqual(self.res["status"], "red")

    def test_numbers_match_reference(self):
        self.assertEqual(self.numbers("red"), sorted(["МП-301", "НЛ-6034", "ТБО-15"]))
        self.assertEqual(self.numbers("yellow"), sorted(["ПВ-08", "МП-322", "ТО-77"]))
        self.assertEqual(self.numbers("grey"), sorted(["АЗ-163", "КП-24"]))
        self.assertEqual(self.numbers("info"), ["НЛ-5521"])

    def test_days_and_normalized_dates(self):
        days, end = self.col["Дней до окончания"], self.col["Дата окончания"]
        expected = {
            "МП-301": (-21, "31.08.2026"), "НЛ-6034": (-21, "31.08.2026"), "ТБО-15": (-7, "14.09.2026"),
            "ПВ-08": (9, "30.09.2026"), "МП-322": (18, "09.10.2026"), "ТО-77": (23, "14.10.2026"),
        }
        for number, (d, e) in expected.items():
            (r, _, _), = self.row(number)
            self.assertEqual((r[days], r[end]), (d, e), number)
        for r, s in zip(self.res["rows"], self.res["row_status"]):
            if s in ("red", "yellow", "green"):
                self.assertRegex(r[end], r"^\d{2}\.\d{2}\.\d{4}$")
                self.assertIsInstance(r[days], int)
            else:
                self.assertIsNone(r[days])

    def test_no_data_and_open_ended(self):
        for number in ("АЗ-163", "КП-24"):
            (r, s, note), = self.row(number)
            self.assertEqual(s, "grey")
            self.assertEqual(note, "нет данных")
        (r, s, note), = self.row("НЛ-5521")
        self.assertEqual((s, note), ("info", "бессрочно"))

    def test_trainer_trap_is_two_green_rows(self):
        found = self.row("ОХ-2025/14")
        self.assertEqual(len(found), 2)
        self.assertEqual({s for _, s, _ in found}, {"green"})
        city = self.col["Город"]
        self.assertEqual(sorted(r[city] for r, _, _ in found), ["Астана", "Караганда"])

    def test_rows_sorted_by_urgency(self):
        order = {"red": 0, "yellow": 1, "green": 2, "grey": 3, "info": 4}
        ranks = [order[s] for s in self.res["row_status"]]
        self.assertEqual(ranks, sorted(ranks))

    def test_schema_fields(self):
        r = self.res
        self.assertEqual(r["title"], "Светофор договоров")
        self.assertRegex(r["updated_at"], r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}")
        self.assertEqual(r["files"], [f"output/{SLUG}/svetofor.xlsx"])
        for row in r["rows"]:
            self.assertEqual(len(row), len(r["columns"]))
        self.assertEqual(len(r["row_status"]), 25)
        self.assertEqual(len(r["row_note"]), 25)
        self.assertNotRegex(json.dumps(r, ensure_ascii=False), "[–—]")

    def test_process_json(self):
        p = json.loads((DEMO / "process.json").read_text(encoding="utf-8"))
        self.assertEqual(p, {
            "title": "Светофор договоров",
            "phrase": "запусти светофор договоров",
            "kind": "script",
            "input": f"input/{SLUG}/",
            "output": f"output/{SLUG}/",
        })


class DemoRunTests(unittest.TestCase):
    """Запуск как со страницы: из корня кабинета в папке с кириллицей и пробелом."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.cab = Path(self._tmp.name) / "Айгуль Н" / "ai-cabinet"
        (self.cab / "tools").mkdir(parents=True)
        shutil.copy(ROOT / "tools" / "read_table.py", self.cab / "tools")
        shutil.copytree(DEMO, self.cab / "processes" / SLUG)
        self.inp = self.cab / "input" / SLUG
        self.inp.mkdir(parents=True)
        self.out = self.cab / "output" / SLUG

    def tearDown(self):
        self._tmp.cleanup()

    def run_demo(self):
        return subprocess.run(
            [sys.executable, str(Path("processes") / SLUG / "run.py")],
            cwd=self.cab, capture_output=True, encoding="utf-8",
            env={**os.environ, "PYTHONUTF8": "1", "PYTHONIOENCODING": "utf-8"},
        )

    def test_writes_result_and_excel(self):
        shutil.copy(CASE_FILE, self.inp)
        p = self.run_demo()
        self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
        res = json.loads((self.out / "result.json").read_text(encoding="utf-8"))
        self.assertEqual(len(res["rows"]), 25)
        self.assertEqual(Counter(res["row_status"])["red"], 3)
        wb = openpyxl.load_workbook(self.out / "svetofor.xlsx", read_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        wb.close()
        self.assertEqual(len(rows), 26)
        self.assertIn("№ договора", rows[0])
        self.assertIn("Статус", rows[0])
        self.assertEqual(sorted(f.name for f in self.out.iterdir()), ["result.json", "svetofor.xlsx"])
        self.assertIn("21.09.2026", p.stdout)

    def test_no_input_file(self):
        p = self.run_demo()
        self.assertEqual(p.returncode, 2)
        self.assertIn(f"input/{SLUG}/reestr_dogovorov.xlsx", p.stdout + p.stderr)
        self.assertRegex(p.stdout + p.stderr, "[А-Яа-я]")
        self.assertNotIn("Traceback", p.stdout + p.stderr)
        self.assertFalse((self.out / "result.json").exists())

    def test_wrong_columns(self):
        wb = openpyxl.Workbook()
        wb.active.append(["Дата", "АЗС", "Литры"])
        wb.active.append(["2026-09-01", "АЗС-01", 100])
        wb.save(self.inp / "reestr_dogovorov.xlsx")
        p = self.run_demo()
        self.assertEqual(p.returncode, 2)
        self.assertIn("Дата окончания", p.stdout + p.stderr)
        self.assertNotIn("Traceback", p.stdout + p.stderr)


if __name__ == "__main__":
    unittest.main()
