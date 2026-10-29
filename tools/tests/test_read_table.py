"""Тесты tools/read_table.py: только стандартная библиотека и openpyxl.

Запуск: python -m unittest discover -s tools/tests -p "test_*.py"
Фикстуры создаются во временной папке с кириллицей и пробелом в имени.
"""

import csv
import json
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path

import openpyxl

TOOLS = Path(__file__).resolve().parent.parent
ROOT = TOOLS.parent
SCRIPT = TOOLS / "read_table.py"
CASE_DATA = ROOT / "input" / "case-azs" / "data"


def run(*args):
    proc = subprocess.run(
        [sys.executable, "-X", "utf8", str(SCRIPT), *map(str, args)],
        capture_output=True,
        encoding="utf-8",
    )
    return proc


def run_json(*args):
    proc = run(*args, "--json")
    if proc.returncode != 0:
        raise AssertionError(f"код {proc.returncode}: {proc.stdout}{proc.stderr}")
    return json.loads(proc.stdout)


class FixtureCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls._tmp.name) / "Папка Айгуль Н"
        cls.dir.mkdir()

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()


class CsvTests(FixtureCase):
    def test_cp1251_semicolon_with_title_row(self):
        p = self.dir / "выгрузка 1с.csv"
        text = (
            "Отчёт по продажам за период 01.09.2026 - 07.09.2026\r\n"
            "\r\n"
            "Дата;АЗС;Товар;Литры\r\n"
            "01.09.2026;АЗС-01;АИ-92;1 200,5\r\n"
            "02.09.2026;АЗС-02;ДТ;800\r\n"
        )
        p.write_bytes(text.encode("cp1251"))
        r = run_json(p)
        self.assertEqual(r["type"], "csv")
        self.assertEqual(r["encoding"], "cp1251")
        self.assertEqual(r["delimiter"], ";")
        self.assertEqual(r["columns"], ["Дата", "АЗС", "Товар", "Литры"])
        self.assertEqual(r["header_row"], 3)
        self.assertEqual(r["data_rows"], 2)
        self.assertEqual(r["rows"][0], ["01.09.2026", "АЗС-01", "АИ-92", "1 200,5"])

    def test_utf8_sig_comma(self):
        p = self.dir / "данные юникод.csv"
        p.write_bytes("﻿Код,Название,Сумма\n1,Мазут,10.5\n2,Бензин,20\n".encode("utf-8"))
        r = run_json(p)
        self.assertEqual(r["encoding"], "utf-8-sig")
        self.assertEqual(r["delimiter"], ",")
        self.assertEqual(r["columns"], ["Код", "Название", "Сумма"])
        self.assertEqual(r["data_rows"], 2)

    def test_plain_utf8_tab_and_rows_limit(self):
        p = self.dir / "tabs.txt"
        lines = ["a\tb\tc"] + [f"{i}\tx{i}\ty{i}" for i in range(30)]
        p.write_text("\n".join(lines), encoding="utf-8")
        r = run_json(p, "--rows", 3)
        self.assertEqual(r["encoding"], "utf-8")
        self.assertEqual(r["delimiter"], "\t")
        self.assertEqual(len(r["rows"]), 3)
        self.assertEqual(r["data_rows"], 30)

    def test_pages_from_and_rows(self):
        """--from N с --rows: строки данных с N-й. Длинный файл читается по 100 строк за раз: страница
        укладывается в вывод команды (около 30 тысяч знаков), а 300 строк сразу нет."""
        p = self.dir / "страницы.csv"
        text = "Колонка не принимает карту, кассир попросил оплатить наличными, а наличных не было. " * 2
        lines = ["№;Текст"] + [f"{i};{text}{i}" for i in range(1, 301)]
        p.write_text("\n".join(lines) + "\n", encoding="utf-8")
        r = run_json(p, "--from", 101, "--rows", 100)
        self.assertEqual(r["data_rows"], 300)
        self.assertEqual(r["first_row"], 101)
        self.assertEqual([row[0] for row in r["rows"]], [str(i) for i in range(101, 201)])
        out = run(p, "--from", 1, "--rows", 100)
        self.assertEqual(out.returncode, 0, out.stderr)
        self.assertIn("Строки 1-100 из 300:", out.stdout)
        self.assertIn(f"  100 | {text}100", out.stdout)
        self.assertIn("Дальше: --from 101 --rows 100.", out.stdout)
        self.assertLess(len(out.stdout), 30000)
        self.assertGreater(len(run(p, "--rows", 300).stdout), 30000)
        out = run(p, "--from", 201, "--rows", 100).stdout
        self.assertIn("Строки 201-300 из 300:", out)
        self.assertIn(f"  300 | {text}300", out)
        self.assertIn("Это последние строки файла.", out)
        self.assertNotIn("Дальше", out)
        self.assertIn("Строк с 301 нет: строк данных 300.", run(p, "--from", 301, "--rows", 100).stdout)
        # Без --from всё как раньше: первые строки, без строки «Дальше».
        out = run(p).stdout
        self.assertIn("Первые строки (10):", out)
        self.assertNotIn("Дальше", out)
        bad = run(p, "--from", 0)
        self.assertEqual(bad.returncode, 2)
        self.assertIn("--from", bad.stderr)

    def test_100k_rows_streamed(self):
        p = self.dir / "большой файл.csv"
        with p.open("w", encoding="cp1251", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Дата", "АЗС", "Литры"])
            for i in range(100_000):
                w.writerow([f"{i % 28 + 1:02d}.09.2026", f"АЗС-{i % 12 + 1:02d}", i])
        t = time.time()
        r = run_json(p)
        self.assertLess(time.time() - t, 20)
        self.assertEqual(r["data_rows"], 100_000)
        self.assertEqual(r["encoding"], "cp1251")

    def test_human_output_in_russian(self):
        p = self.dir / "человек.csv"
        p.write_bytes("Код;Имя\n1;Ёлка\n".encode("cp1251"))
        out = run(p)
        self.assertEqual(out.returncode, 0)
        self.assertIn("cp1251", out.stdout)
        self.assertIn("Разделитель", out.stdout)
        self.assertIn("Строка шапки", out.stdout)
        # Шапка в первой строке: строки «Над шапкой» нет.
        self.assertNotIn("Над шапкой", out.stdout)
        self.assertEqual(run_json(p)["above_header"], [])


class XlsxTests(FixtureCase):
    def test_two_level_merged_header(self):
        p = self.dir / "сводная шапка.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Свод"
        ws["A1"] = "Сводка по сети за сентябрь"
        ws["A2"] = "АЗС"
        ws.merge_cells("A2:A3")
        ws["B2"] = "Продажи"
        ws.merge_cells("B2:C2")
        ws["D2"] = "Остаток"
        ws.merge_cells("D2:E2")
        for col, v in zip("BCDE", ["Бензин", "Дизель", "Бензин", "Дизель"]):
            ws[f"{col}3"] = v
        ws.append(["АЗС-01", 10, 20, 30, 40])
        ws.append(["АЗС-02", 11, 21, 31, 41])
        wb.save(p)
        r = run_json(p)
        self.assertEqual(r["type"], "xlsx")
        self.assertEqual(
            r["columns"],
            ["АЗС", "Продажи / Бензин", "Продажи / Дизель", "Остаток / Бензин", "Остаток / Дизель"],
        )
        self.assertEqual(r["header_row"], 2)
        self.assertEqual(r["header_rows"], 2)
        self.assertEqual(r["data_rows"], 2)
        self.assertEqual(r["rows"][0], ["АЗС-01", "10", "20", "30", "40"])

    def test_rows_above_header_shown(self):
        """Заголовок отчёта, поставщик и период над шапкой: Claude видит их, программа берёт из above_header."""
        p = self.dir / "сверка с заголовком.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Сверка"
        ws["A1"] = "Сверка с ТОО «Ромашка» за август 2026"
        ws.merge_cells("A1:E1")
        ws["A2"] = "Поставщик:"
        ws["B2"] = "ТОО «Ромашка»"
        ws.append([])
        ws.append(["Документ", "Дата", "Сумма у нас", "Сумма у поставщика", "Комментарий"])
        ws.append(["№ 41", "05.08.2026", 6195000, None, "нет у поставщика"])
        ws.append(["№ 42", "12.08.2026", 1250000, 1200000, None])
        wb.save(p)
        r = run_json(p)
        self.assertEqual(r["header_row"], 4)
        self.assertEqual(r["above_header"], [["Сверка с ТОО «Ромашка» за август 2026"], ["Поставщик:", "ТОО «Ромашка»"]])
        out = run(p)
        self.assertEqual(out.returncode, 0, out.stderr)
        self.assertIn("Над шапкой: Сверка с ТОО «Ромашка» за август 2026; Поставщик: | ТОО «Ромашка»", out.stdout)
        lines = out.stdout.splitlines()
        self.assertLess(lines.index(next(x for x in lines if x.startswith("Над шапкой"))),
                        lines.index(next(x for x in lines if x.startswith("Строка шапки"))))
        # --rows ограничивает только первые строки данных: строки над шапкой те же.
        self.assertEqual(len(run_json(p, "--rows", 1)["above_header"]), 2)
        # Другой лист без строк над шапкой: строки «Над шапкой» нет.
        wb = openpyxl.load_workbook(p)
        ws2 = wb.create_sheet("Итог")
        ws2.append(["Документ", "Сумма"])
        ws2.append(["№ 41", 1])
        wb.save(p)
        self.assertNotIn("Над шапкой", run(p, "--sheet", "Итог").stdout)
        self.assertIn("Над шапкой: Сверка", run(p, "--sheet", "Сверка").stdout)

    def test_rows_above_header_cap_and_long_cells(self):
        """Над шапкой больше 15 строк: первые 15 и «... и ещё N»; длинная ячейка обрезается. В JSON всё целиком."""
        p = self.dir / "много строк над шапкой.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        title = "Акт сверки " + "очень длинный заголовок " * 20
        ws.append([title])
        for i in range(2, 18):
            ws.append([f"строка {i}"])
        ws.append(["Документ", "Дата", "Сумма"])
        ws.append(["№ 1", "01.08.2026", 10])
        wb.save(p)
        r = run_json(p)
        self.assertEqual(r["header_row"], 18)
        self.assertEqual(len(r["above_header"]), 17)
        self.assertEqual(r["above_header"][0], [title.strip()])
        out = run(p).stdout
        line = next(x for x in out.splitlines() if x.startswith("Над шапкой: "))
        self.assertIn("строка 15", line)
        self.assertNotIn("строка 16", line)
        self.assertTrue(line.endswith("... и ещё 2 строк над шапкой"), line)
        first = line[len("Над шапкой: "):].split("; ")[0]
        self.assertEqual(first, title.strip()[:200] + "...")

    def test_three_sheets(self):
        p = self.dir / "три листа.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Астана"
        wb.active.append(["Договор", "Сумма"])
        for i in range(3):
            wb.active.append([f"Д-{i}", i])
        ws2 = wb.create_sheet("Караганда")
        ws2.append(["Договор", "Сумма"])
        for i in range(5):
            ws2.append([f"К-{i}", i])
        ws3 = wb.create_sheet("Справочник")
        ws3.append(["Код", "Название"])
        ws3.append([1, "Бензин"])
        wb.save(p)
        r = run_json(p)
        self.assertEqual([s["name"] for s in r["sheets"]], ["Астана", "Караганда", "Справочник"])
        self.assertEqual([s["rows"] for s in r["sheets"]], [4, 6, 2])
        self.assertEqual(r["sheet"]["name"], "Астана")
        r2 = run_json(p, "--sheet", 2)
        self.assertEqual(r2["sheet"]["name"], "Караганда")
        self.assertEqual(r2["data_rows"], 5)
        r3 = run_json(p, "--sheet", "Справочник")
        self.assertEqual(r3["columns"], ["Код", "Название"])
        self.assertEqual(r3["rows"], [["1", "Бензин"]])

    def test_vertical_merge_over_data_is_single_level_header(self):
        p = self.dir / "merge вниз.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["АЗС", "Показатель", "Знач"])
        ws.append([None, "a", 1])
        ws.append([None, "b", 2])
        ws.append([None, "c", 3])
        ws.append(["АЗС-02", "d", 4])
        ws.merge_cells("A1:A4")
        wb.save(p)
        r = run_json(p, "--rows", 20)
        self.assertEqual(r["columns"], ["АЗС", "Показатель", "Знач"])
        self.assertEqual(r["header_rows"], 1)
        self.assertEqual(r["data_rows"], 4)
        self.assertEqual(r["rows"][0], ["АЗС", "a", "1"])
        self.assertEqual(r["rows"][2], ["АЗС", "c", "3"])
        self.assertEqual(r["rows"][3], ["АЗС-02", "d", "4"])

    def test_three_level_header(self):
        p = self.dir / "три уровня.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Группа"
        ws.merge_cells("A1:C1")
        ws["A2"] = "Подгр1"
        ws.merge_cells("A2:B2")
        ws["C2"] = "Подгр2"
        for col, v in zip("ABC", ["x", "y", "z"]):
            ws[f"{col}3"] = v
        ws.append([1, 2, 3])
        ws.append([4, 5, 6])
        wb.save(p)
        r = run_json(p)
        self.assertEqual(r["columns"], ["Группа / Подгр1 / x", "Группа / Подгр1 / y", "Группа / Подгр2 / z"])
        self.assertEqual(r["header_rows"], 3)
        self.assertEqual(r["header_row"], 1)
        self.assertEqual(r["data_rows"], 2)
        self.assertEqual(r["rows"], [["1", "2", "3"], ["4", "5", "6"]])

    def test_unknown_sheet(self):
        p = self.dir / "лист для ошибки.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Астана"
        wb.active.append(["a", "b"])
        wb.save(p)
        out = run(p, "--sheet", "Нет такого")
        self.assertEqual(out.returncode, 2)
        self.assertIn("Астана", out.stderr)


    def test_sheet_as_number_from_program(self):
        # Программа процесса передаёт номер листа числом: read_table.analyze(путь, 2, ...).
        sys.path.insert(0, str(TOOLS))
        try:
            import read_table
        finally:
            sys.path.remove(str(TOOLS))
        p = self.dir / "листы числом.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Астана"
        wb.active.append(["Договор", "Сумма"])
        wb.active.append(["А-1", 1])
        ws2 = wb.create_sheet("Караганда")
        ws2.append(["Договор", "Сумма"])
        ws2.append(["К-1", 2])
        ws2.append(["К-2", 3])
        wb.save(p)
        r = read_table.analyze(p, 2, 10)
        self.assertEqual(r["sheet"]["name"], "Караганда")
        self.assertEqual(r["data_rows"], 2)
        self.assertEqual(read_table.analyze(p, 1.0, 10)["sheet"]["name"], "Астана")
        self.assertEqual(read_table.analyze(p, " 2 ", 10)["sheet"]["name"], "Караганда")
        for bad in (3, 0, -1, 1.5, True, [1]):
            with self.assertRaises(read_table.UserError) as ctx:
                read_table.analyze(p, bad, 10)
            msg = str(ctx.exception)
            self.assertIn("Астана", msg, bad)
            self.assertRegex(msg, "[А-Яа-я]")
        with self.assertRaises(read_table.UserError) as ctx:
            read_table.analyze(p, 3, 10)
        self.assertIn("номера с 1", str(ctx.exception))


class TotalsTests(FixtureCase):
    """Итоги по всем строкам (--sum, --by, --where): ими Claude делает эталон расчёта независимо от программы."""

    def write(self):
        p = self.dir / "продажи итоги.csv"
        text = (
            "Отчёт\r\n\r\n"
            "Дата;АЗС;Операция;Литры;Сумма, тг\r\n"
            "07.09.2026;АЗС-01;Реализация;1 000,5;245 122,50\r\n"
            "08.09.2026;АЗС-02;Реализация;500;122 500,00\r\n"
            "14.09.2026;АЗС-01;Реализация;200;49 000,00\r\n"
            "14.09.2026;АЗС-01;Возврат;-50;-12 250,00\r\n"
            "15.09.2026;АЗС-02;Реализация;;0\r\n"
            "16.09.2026;АЗС-02;Реализация;н/д;0\r\n"
        )
        p.write_bytes(text.replace("\u00a0", " ").encode("cp1251"))
        return p

    def test_sum_by_where(self):
        p = self.write()
        t = run_json(p, "--by", "азс", "--sum", "Литры", "--sum", "сумма", "--where", "Операция=реализация")
        self.assertEqual(t["by"], ["АЗС"])
        self.assertEqual(t["sums"], ["Литры", "Сумма, тг"])
        self.assertEqual(t["total"]["rows"], 5)
        groups = {g["key"][0]: g for g in t["groups"]}
        self.assertEqual(list(groups), ["АЗС-01", "АЗС-02"])
        self.assertAlmostEqual(groups["АЗС-01"]["sums"][0], 1200.5)
        self.assertEqual(groups["АЗС-02"]["rows"], 3)
        self.assertAlmostEqual(t["total"]["sums"][1], 416622.5)
        self.assertEqual(t["total"]["empty"], [1, 0])
        self.assertEqual(t["total"]["bad"], [[6], []])  # номер строки данных в файле

    def test_date_range_and_or_of_same_column(self):
        p = self.write()
        t = run_json(p, "--sum", "Литры", "--where", "Дата=14.09.2026..20.09.2026")
        self.assertEqual(t["total"]["rows"], 4)
        self.assertAlmostEqual(t["total"]["sums"][0], 150)
        t = run_json(p, "--sum", "4", "--where", "АЗС=АЗС-01", "--where", "АЗС=АЗС-02", "--where", "Операция=Возврат")
        self.assertEqual(t["total"]["rows"], 1)
        self.assertAlmostEqual(t["total"]["sums"][0], -50)

    def test_human_output_and_errors(self):
        p = self.write()
        out = run(p, "--by", "АЗС", "--by", "Операция", "--sum", "Литры")
        self.assertEqual(out.returncode, 0, out.stderr)
        self.assertIn("АЗС | Операция | строк | Литры", out.stdout)
        self.assertIn("АЗС-01 | Реализация | 2 | 1200.5", out.stdout)
        self.assertIn("Всего |  | 6 | 1650.5", out.stdout)
        self.assertIn("«Литры»: не числа в 1 строках данных (6), в сумму не вошли.", out.stdout)
        bad = run(p, "--sum", "Цена")
        self.assertEqual(bad.returncode, 2)
        self.assertIn("Столбец «Цена» не найден", bad.stderr)
        both = run(p, "--sum", "а")
        self.assertEqual(both.returncode, 2)
        self.assertIn("несколько столбцов, укажите номер", both.stderr)
        rng = run(p, "--sum", "Литры", "--where", "Дата=вчера..сегодня")
        self.assertEqual(rng.returncode, 2)
        self.assertIn("две даты", rng.stderr)

    def test_1c_datetime_in_range_and_unreadable_dates_reported(self):
        p = self.dir / "1с даты.csv"
        text = (
            "Дата;АЗС;Литры\r\n"
            "14.09.2026 0:00:00;АЗС-01 ;100\r\n"
            "15.09.2026 13:45:10;азс-01;50\r\n"
            "2026-09-16 10:30:00;АЗС-02;25\r\n"
            "вчера;АЗС-02;7\r\n"
            ";АЗС-02;3\r\n"
            "21.09.2026 0:00:00;АЗС-02;1\r\n"
        )
        p.write_bytes(text.encode("cp1251"))
        t = run_json(p, "--by", "АЗС", "--sum", "Литры", "--where", "Дата=14.09.2026..20.09.2026")
        self.assertEqual(t["total"]["rows"], 3)
        self.assertAlmostEqual(t["total"]["sums"][0], 175)
        groups = {g["key"][0]: g for g in t["groups"]}
        self.assertEqual(list(groups), ["АЗС-01", "АЗС-02"])  # «АЗС-01 » и «азс-01» одна группа
        self.assertEqual(groups["АЗС-01"]["rows"], 2)
        self.assertEqual(t["unreadable"], [{"column": "Дата", "kind": "дата", "rows": 2}])
        out = run(p, "--sum", "Литры", "--where", "Дата=14.09.2026..20.09.2026")
        self.assertIn("Внимание: «Дата»: дата не читается в 2 строках, в отбор не вошли.", out.stdout)
        ok = run(p, "--sum", "Литры")
        self.assertNotIn("не читается", ok.stdout)

    def test_number_digits_only(self):
        import importlib.util
        spec = importlib.util.spec_from_file_location("read_table_num", SCRIPT)
        rt = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(rt)
        self.assertEqual(rt.to_number("1 234,5"), 1234.5)
        for bad in ("nan", "inf", "1_000", "1e5", "Infinity"):
            self.assertIsNone(rt.to_number(bad), bad)

    def test_case_sales_week(self):
        t = run_json(CASE_DATA / "sales_azs_1c.csv", "--sum", "Количество", "--where", "Дата=14.09.2026..20.09.2026",
                     "--where", "Операция=Реализация")
        self.assertEqual(t["total"]["rows"], 652)
        self.assertAlmostEqual(t["total"]["sums"][0], 862935.8, places=3)


class ErrorTests(FixtureCase):
    def test_missing_file(self):
        out = run(self.dir / "нет.csv")
        self.assertEqual(out.returncode, 2)
        self.assertIn("не найден", out.stderr)

    def test_old_xls(self):
        p = self.dir / "старый.xls"
        p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64)
        out = run(p)
        self.assertEqual(out.returncode, 2)
        self.assertIn("сохраните как .xlsx", out.stderr)

    def test_password_protected_xlsx(self):
        p = self.dir / "закрыт.xlsx"
        p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64)
        out = run(p)
        self.assertEqual(out.returncode, 2)
        self.assertIn("пароль", out.stderr)

    def test_unsupported_type(self):
        p = self.dir / "doc.pdf"
        p.write_bytes(b"%PDF-1.4")
        out = run(p)
        self.assertEqual(out.returncode, 2)
        self.assertIn("не поддерживается", out.stderr)


class CaseFilesSmoke(unittest.TestCase):
    def test_every_packed_case_table_reads(self):
        files = sorted(
            f for f in (ROOT / "input" / "case-azs").rglob("*") if f.suffix.lower() in (".csv", ".xlsx")
        )
        self.assertGreater(len(files), 10, "кейс не упакован: запустите pack_case.sh")
        for f in files:
            with self.subTest(file=str(f.relative_to(ROOT))):
                r = run_json(f)
                self.assertTrue(r["columns"], "шапка не найдена")
                self.assertGreater(r["data_rows"], 0)

    def test_registry_all_sheets(self):
        f = CASE_DATA / "reestr_dogovorov.xlsx"
        r = run_json(f)
        self.assertGreaterEqual(len(r["sheets"]), 3)
        for i in range(1, len(r["sheets"]) + 1):
            run_json(f, "--sheet", i)


if __name__ == "__main__":
    unittest.main()


class BadColors(unittest.TestCase):
    """Выгрузка с цветом не по стандарту (бывает у 1С и Google Таблиц): файл всё равно читается."""

    def test_xlsx_with_invalid_rgb_is_read(self):
        import shutil
        import zipfile
        from openpyxl.styles import PatternFill
        with tempfile.TemporaryDirectory() as d:
            good = Path(d) / "good.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["АЗС", "Литры"])
            ws.append(["АЗС-01", 10])
            ws["A1"].fill = PatternFill("solid", fgColor="FFFF0000")
            wb.save(good)
            bad = Path(d) / "Выгрузка с цветом.xlsx"
            with zipfile.ZipFile(good) as zin, zipfile.ZipFile(bad, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename == "xl/styles.xml":
                        data = data.replace(b'rgb="FFFF0000"', b'rgb="#FF00"')
                    zout.writestr(item, data)
            r = subprocess.run([sys.executable, str(SCRIPT), str(bad), "--json"], capture_output=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            self.assertIn("АЗС-01", r.stdout)
