"""Образцы программы процесса из скилла add-process: основа, словарь, шаблон Word (programma.md),
два документа Word, сверка договора с типовой формой (programma-dva-dokumenta.md).

Каждый образец собирается в программу так, как велит справка: основа, блок, пример main. Программа
запускается как программа процесса, её result.json проверяет node tools/processes.mjs list (та же схема,
что у карточки), а сверку со скиллом node tools/processes.mjs compare.

Запуск: python -m unittest discover -s tools/tests -p "test_*.py"
"""

import copy
import importlib.util
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

TOOLS = Path(__file__).resolve().parent.parent
ROOT = TOOLS.parent
GUIDE = ROOT / ".claude" / "skills" / "add-process" / "programma.md"
GUIDE_DVA = ROOT / ".claude" / "skills" / "add-process" / "programma-dva-dokumenta.md"
CASE = ROOT / "input" / "case-azs" / "data"
NODE = shutil.which("node")

try:
    import openpyxl  # noqa: F401
    import docx  # noqa: F401
    HAVE_LIBS = True
except ImportError:
    HAVE_LIBS = False


def blocks():
    found = re.findall(r"```python\n(.*?)```", GUIDE.read_text(encoding="utf-8"), re.S)
    found_dva = re.findall(r"```python\n(.*?)```", GUIDE_DVA.read_text(encoding="utf-8"), re.S)
    assert len(found) == 5, f"в programma.md нужны 5 образцов, найдено {len(found)}"
    assert len(found_dva) == 2, f"в programma-dva-dokumenta.md нужны 2 образца, найдено {len(found_dva)}"
    named = {}
    for b in found + found_dva:
        if b.startswith('"""<Название процесса>'):
            named["base"] = b
        elif b.startswith("# --- Словарь"):
            named["slovar"] = b
        elif b.startswith("# --- Шаблон Word"):
            named["word"] = b
        elif "slovar = load_slovar()" in b and b.startswith("def main():"):
            named["main_slovar"] = b
        elif "fill_docx(" in b and b.startswith("def main():"):
            named["main_word"] = b
        elif b.startswith("# --- Два документа Word"):
            named["dva"] = b
        elif "two_documents(" in b and b.startswith("def main():"):
            named["main_dva"] = b
    assert len(named) == 7, f"в справке нужны 7 образцов, найдено {sorted(named)}"
    return named


def compose(*extra):
    """Основа без своего main, затем блоки и пример main перед запуском."""
    b = blocks()
    base = b["base"].replace("<имя>", "proba").replace("<Название процесса>", "Проба")
    if not extra:
        return base
    head, sep, tail = base.partition("\ndef main():")
    assert sep, "в основе нет def main()"
    run_part = tail[tail.index("\nif __name__"):]
    return head + "\n" + "\n\n".join(b[name] for name in extra) + "\n" + run_part


class Cabinet:
    def __init__(self, code):
        self.tmp = tempfile.TemporaryDirectory()
        self.cab = Path(self.tmp.name) / "Айгуль Н" / "ai-cabinet"
        shutil.copytree(TOOLS, self.cab / "tools", ignore=shutil.ignore_patterns("tests", "__pycache__"))
        self.proc = self.cab / "processes" / "proba"
        self.proc.mkdir(parents=True)
        (self.proc / "run.py").write_text(code, encoding="utf-8")
        (self.proc / "process.json").write_text(json.dumps(
            {"title": "Проба", "phrase": "запусти пробу", "kind": "script"}, ensure_ascii=False), encoding="utf-8")
        self.input = self.cab / "input" / "proba"
        self.input.mkdir(parents=True)
        self.out = self.cab / "output" / "proba"

    def close(self):
        self.tmp.cleanup()

    def run(self):
        return subprocess.run([sys.executable, str(self.proc / "run.py")], cwd=self.cab,
                              capture_output=True, encoding="utf-8")

    def result(self):
        return json.loads((self.out / "result.json").read_text(encoding="utf-8"))

    def module(self):
        """Программа как модуль: её функции для проверки по отдельности."""
        spec = importlib.util.spec_from_file_location(f"run_proba_{id(self)}", self.proc / "run.py")
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod

    def node(self, *args):
        r = subprocess.run([NODE, str(self.cab / "tools" / "processes.mjs"), *args], cwd=self.cab,
                           capture_output=True, encoding="utf-8")
        return r.returncode, r.stdout + r.stderr


def manual_rows(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()


@unittest.skipUnless(HAVE_LIBS, "нет openpyxl или python-docx")
class Base(unittest.TestCase):
    def run_base(self, rows, limit=None):
        code = compose()
        if limit is not None:
            code, n = re.subn(r"^LIMIT = .*$", f"LIMIT = {limit}", code, flags=re.M)
            self.assertEqual(n, 1, "в образце нет строки LIMIT = ...")
        c = Cabinet(code)
        self.addCleanup(c.close)
        lines = ["АЗС;Литры"] + [f"АЗС-{i:02d};{i * 10}" for i in range(1, rows + 1)]
        (c.input / "sales.csv").write_text("\n".join(lines) + "\n", encoding="utf-8")
        proc = c.run()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        return c, c.result()

    def test_all_rows_no_warning(self):
        c, r = self.run_base(5)
        self.assertNotIn("Внимание", r["summary"])
        self.assertIn("sales.csv", r["summary"])
        self.assertIn("Разобрать вручную: 0", r["summary"])

    def test_truncated_rows_warned_in_summary(self):
        _, r = self.run_base(5, limit=3)
        self.assertIn("Внимание: в файле 5 строк, посчитаны первые 3.", r["summary"])

    def test_manual_file_written_even_empty(self):
        c, r = self.run_base(2)
        self.assertEqual(r["files"], ["output/proba/razobrat-vruchnuyu.xlsx"])
        rows = manual_rows(c.out / "razobrat-vruchnuyu.xlsx")
        self.assertEqual(rows[0], ["АЗС", "Литры", "Почему"])
        self.assertEqual(rows[1][0], "Всё разобрано программой.")
        self.assertIn("sverka", r)

    def test_number_digits_only(self):
        c = Cabinet(compose())
        self.addCleanup(c.close)
        m = c.module()
        self.assertEqual(m.number("1 234,5"), 1234.5)
        self.assertEqual(m.number("374\u00a0433,50"), 374433.5)
        self.assertEqual(m.number("-12"), -12.0)
        for bad in ("nan", "inf", "-inf", "1_000", "1e5", "", "н/д", "1.2.3"):
            self.assertIsNone(m.number(bad), bad)

    def test_no_input_clear_stop(self):
        c = Cabinet(compose())
        self.addCleanup(c.close)
        proc = c.run()
        self.assertEqual(proc.returncode, 2)
        self.assertIn("В папке input/proba нет файла", proc.stderr)
        self.assertNotIn("Traceback", proc.stderr)


SLOVAR = {
    "качество топлива": ["качество топлива", "топливо", "аи 92", "дт", "дизель", "жанармай", "недолив"],
    "обслуживание и персонал": ["оператор", "кассир", "нахамил", "нагрубил", "очередь", "вежлив"],
    "оплата и касса": ["чек", "списали", "терминал", "оплата", "касса"],
    "минимаркет и кафе": ["кофе", "кафе", "выпечка", "хот-дог"],
    "программа лояльности": ["бонус", "карта лояльности", "акция"],
    "чистота и санузлы": ["туалет", "санузел", "санузл", "грязно", "дәретхана"],
    "оборудование": ["колонка", "колонок", "пистолет", "пылесос", "табло", "подкачка"],
    "благодарность": ["спасибо", "благодар", "рахмет"],
}


def guide_slovar_example():
    """Пример slovar.json из раздела «Словарь» справки: в нём слова срочности по умолчанию."""
    text = GUIDE.read_text(encoding="utf-8")
    part = text[text.index("## Словарь"):]
    return json.loads(re.search(r"```json\n(.*?)```", part, re.S).group(1))


@unittest.skipUnless(HAVE_LIBS, "нет openpyxl или python-docx")
class Slovar(unittest.TestCase):
    def setUp(self):
        self.c = Cabinet(compose("slovar", "main_slovar"))
        self.addCleanup(self.c.close)
        (self.c.proc / "slovar.json").write_text(json.dumps(SLOVAR, ensure_ascii=False), encoding="utf-8")

    def with_urgent(self, extra=None):
        """Словарь теста и слова срочности по умолчанию из справки, как их берёт Claude."""
        m = self.c.module()
        urgent = guide_slovar_example()[m.URGENT]
        (self.c.proc / "slovar.json").write_text(json.dumps(dict(SLOVAR, **{m.URGENT: urgent}, **(extra or {})),
                                                            ensure_ascii=False), encoding="utf-8")
        return m

    def run_case(self, name):
        for old in self.c.input.glob("*"):
            old.unlink()
        shutil.copy(CASE / name, self.c.input / name)
        proc = self.c.run()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        return self.c.result()

    def test_word_rule(self):
        m = self.c.module()
        s = m.load_slovar()
        cls = lambda t: m.classify(t, s)[0]
        self.assertEqual(m.stem("бонусы"), "бонус")
        self.assertEqual(m.stem("касса"), "касс")
        self.assertEqual(m.stem("очередь"), "очеред")
        self.assertEqual(m.stem("мыло"), "мыл")
        self.assertEqual(m.stem("чек"), "чек")
        self.assertEqual(cls("Почему бонусов так мало?"), "программа лояльности")
        self.assertEqual(cls("Заправился ДТ, машина не заводится"), "качество топлива")
        self.assertEqual(cls("После ДТП стоял у вас на заправке"), m.MANUAL)  # «дт» только целым словом
        self.assertEqual(cls("АИ-92 явно хуже, чем у соседей"), "качество топлива")  # фраза «аи 92»
        self.assertEqual(cls("ЖАНАРМАЙ сапасы нашар"), "качество топлива")  # регистр, казахский
        self.assertEqual(cls("Дәретханада қағаз жоқ"), "чистота и санузлы")
        self.assertEqual(cls("Ещё раз про бонусы"), "программа лояльности")
        # Найденные слова словаря: по ним человек видит, почему строка в теме.
        self.assertEqual(m.classify("Почему бонусов мало и акция не работает?", s),
                         ("программа лояльности", "", ["бонус", "акция"]))

    def test_manual_reasons(self):
        m = self.c.module()
        s = m.load_slovar()
        self.assertEqual(m.classify("", s), (m.MANUAL, "пустой текст", []))
        self.assertEqual(m.classify("Можно ли оставить машину на ночь?", s), (m.MANUAL, "нет слов из словаря", []))
        cat, why, found = m.classify("Кофе холодный, а оператор нагрубил", s)  # кафе 1, персонал 2
        self.assertEqual((cat, why, found), ("обслуживание и персонал", "", ["оператор", "нагрубил"]))

    def test_tie_goes_to_topic_written_first(self):
        """Поровну у двух тем: тема слова, которое в тексте раньше, с пометкой «две темы», не «разобрать вручную»."""
        m = self.c.module()
        s = m.load_slovar()
        self.assertEqual(m.classify("Кофе холодный, а оператор молчал", s),
                         ("минимаркет и кафе", "две темы: минимаркет и кафе, обслуживание и персонал", ["кофе"]))
        self.assertEqual(m.classify("Оператор молчал, а кофе холодный", s)[:2],
                         ("обслуживание и персонал", "две темы: обслуживание и персонал, минимаркет и кафе"))
        cat, why, _ = m.classify("Туалет грязный, кофе холодный, оператор молчал", s)
        self.assertEqual(cat, "чистота и санузлы")  # чистота 2 слова: не ничья
        self.assertFalse(why.startswith(m.TIES))
        cat, why, _ = m.classify("Колонка сломана, кофе холодный, кассир молчал", s)
        self.assertEqual((cat, why.split(":")[0]), ("оборудование", "несколько тем"))
        # Одно слово в двух темах (одно место в тексте): темы по алфавиту, тоже с пометкой «две темы».
        both = dict(s, **{"оборудование": s["оборудование"] + [("экран", m.roots("экран"))],
                          "оплата и касса": s["оплата и касса"] + [("экран", m.roots("экран"))]})
        self.assertEqual(m.classify("Экран погас", both)[:2],
                         ("оборудование", "две темы: оборудование, оплата и касса"))

    def test_stem_word_families(self):
        """Основа у слова словаря и у слова текста одна: слово в обычной форме находит другие формы."""
        m = self.c.module()
        families = [
            "зарядка зарядки зарядные зарядку", "колонка колонки колонке колонок колонками",
            "бонус бонусы бонусов бонусами", "касса кассу кассе кассой", "грязь грязи грязный грязно грязная",
            "топливо топлива топливом топливная топливной", "наличные наличными наличных наличка",
            "благодарность благодарю благодарим", "очередь очереди очередью очередей",
            "оплата оплату оплатой оплатить", "сотрудник сотрудника сотрудники сотрудникам", "мыло мыла мылом",
            "кассир кассира кассиру кассиры", "балл баллы баллов балла",
        ]
        # Слово словаря в обычной форме (первое) находит остальные: казахские окончания, суффикс, короткое слово.
        from_base = ["командировка командировки командировочные", "дәретхана дәретханада дәретханасы",
                     "жанармай жанармайдың жанармайды", "ұпай ұпайлар ұпайлары", "газ газом газовой газа",
                     "запах запаха запахом", "пол пола полу полы", "сок сока соки",
                     "қызмет қызметі қызметке қызметтің", "кезек кезекте кезекке", "утечка утечку утечки",
                     "иіс иісі", "бензин бензина бензиновый"]
        for fam in families + from_base:
            words = fam.split()
            pairs = [(a, b) for a in words for b in words] if fam in families else [(words[0], b) for b in words]
            for entry, text in pairs:
                with self.subTest(entry=entry, text=text):
                    slovar = {"тема": [(entry, m.roots(entry))]}
                    self.assertEqual(m.classify(text, slovar)[0], "тема")
        self.assertEqual(m.stem("зарядные"), m.stem("зарядка"))
        self.assertEqual(m.stem("топливная"), m.stem("топливо"))
        for short in ("газ", "чек", "сок", "дт"):  # до трёх букв слово не укорачивается
            self.assertEqual(m.stem(short), short)
        self.assertEqual(m.stem("машина"), "машин")  # корень на -н не режется до гласной
        # Другое слово с тем же началом не находится: после слова словаря должно идти окончание.
        others = [("кофе", "Купил сок и печенье"), ("запах", "Заправляюсь только у вас"), ("дт", "После ДТП"),
                  ("балл", "баллон"), ("газ", "газету"), ("газ", "газон"), ("газ", "газель"),
                  ("газ", "газированная"), ("колонка", "колонну"), ("колонка", "колонна"), ("пол", "полицию"),
                  ("пол", "полный"), ("пол", "полка"), ("пол", "получил"), ("сок", "сокол"), ("сок", "сокращение"),
                  ("қызмет", "қызметкер"), ("касса", "кассир"), ("балл", "баллада"), ("бонус", "бонусник")]
        for entry, text in others:
            with self.subTest(entry=entry, text=text):
                self.assertEqual(m.classify(text, {"тема": [(entry, m.roots(entry))]})[0], m.MANUAL)

    def test_broken_slovar_clear_stop(self):
        (self.c.proc / "slovar.json").write_text('{"тема": ["слово",]}', encoding="utf-8")
        shutil.copy(CASE / "obrashcheniya.xlsx", self.c.input / "obrashcheniya.xlsx")
        proc = self.c.run()
        self.assertEqual(proc.returncode, 2, proc.stdout + proc.stderr)
        self.assertIn("Словарь processes/proba/slovar.json испорчен (строка 1)", proc.stderr)
        self.assertNotIn("Traceback", proc.stderr)
        (self.c.proc / "slovar.json").unlink()
        proc = self.c.run()
        self.assertEqual(proc.returncode, 2)
        self.assertIn("Нет словаря processes/proba/slovar.json", proc.stderr)

    def test_case_file_card_manual_list_and_check(self):
        r = self.run_case("obrashcheniya.xlsx")
        self.assertEqual(len(r["rows"]), 80)
        self.assertEqual(r["status"], "green")  # в словаре нет раздела срочности: строк «срочно» нет
        self.assertFalse(r["summary"].startswith("Срочно"))
        manual = [i for i, n in enumerate(r["row_note"]) if n == "разобрать вручную"]
        self.assertTrue(manual, "на файле кейса есть строки без слов из словаря")
        self.assertEqual(manual, list(range(len(manual))), "строки «разобрать вручную» идут первыми")
        self.assertTrue(all(r["row_status"][i] == "info" for i in manual))
        rest = list(zip(r["row_status"], r["row_note"]))[len(manual):]
        ties = [n for s, n in rest if s == "info"]
        self.assertTrue(ties, "на файле кейса есть строки с двумя темами")
        self.assertTrue(all(n.startswith(("две темы: ", "несколько тем: ")) for n in ties), ties)
        self.assertTrue(all(n is None for s, n in rest if s is None))
        self.assertIn(f"Две темы: {len(ties)}", r["summary"])
        self.assertIn(f"Разобрать вручную: {len(manual)} (файл razobrat-vruchnuyu.xlsx)", r["summary"])
        self.assertIn("output/proba/razobrat-vruchnuyu.xlsx", r["files"])
        self.assertEqual(len(manual_rows(self.c.out / "razobrat-vruchnuyu.xlsx")), len(manual) + 1)
        rows = r["sverka"]["по строкам"]
        self.assertEqual(len(rows), 80)
        self.assertEqual(rows["38"], "повтор")  # повтор обращения 22 той же АЗС
        self.assertEqual(rows["79"], "повтор")
        self.assertEqual(r["sverka"]["итоги"], {"обращений": 80, "повторов": 2})
        self.assertEqual(r["sverka"]["срочно"], [])
        self.assertEqual(sum(r["sverka"]["по категориям"].values()), 80 - 2 - len(manual))
        # В Excel результата у каждой строки слова словаря, по которым она попала в тему.
        table = sheet(self.c.out / "obrashcheniya-po-temam.xlsx", "Обращения")
        self.assertEqual(table[0], ["№", "АЗС", "Тема", "Текст", "Найдено по слову", "Пометка"])
        by_num = {row[0]: row for row in table[1:]}
        self.assertEqual(by_num["5"][2:5], ["программа лояльности", "Почему бонусы сгорают так быстро? Не успел потратить.",
                                            "бонус"])
        self.assertEqual(by_num["38"][4:], [None, "повтор, посчитан один раз"])
        self.assertEqual(by_num[r["rows"][0][0]][5], "нет слов из словаря")
        if not NODE:
            return
        code, out = self.c.node("list")
        self.assertEqual(code, 0, out)
        self.assertIn("Результат:", out)
        self.assertNotIn("Ошибка", out)
        etalon = {k: r["sverka"][k] for k in ("итоги", "по категориям")}
        etalon["по строкам"] = dict(rows)
        first = r["rows"][0][0]  # строка «разобрать вручную»: у скилла своя тема
        etalon["по строкам"][first] = "другое"
        etalon["по категориям"] = dict(etalon["по категориям"], **{"другое": 1})
        (self.c.out / "etalon-skilla.json").write_text(json.dumps(etalon, ensure_ascii=False), encoding="utf-8")
        code, out = self.c.node("compare", "proba")
        self.assertEqual(code, 0, out)
        self.assertIn("кроме", out)

    def test_urgent_words_default_in_guide(self):
        """Раздел срочности в примере словаря справки: имя раздела как в программе, опасность по-русски и по-казахски."""
        m = self.c.module()
        example = guide_slovar_example()
        self.assertIn(m.URGENT, example)
        urgent = example[m.URGENT]
        for word in ("запах", "пахнет", "утечка", "разлив", "дым", "искры", "пожар", "огонь", "взрыв", "травма",
                     "отравление", "драка", "угроза", "иіс", "өрт", "түтін", "жарылыс", "жарақат"):
            self.assertIn(word, urgent)
        # Срочность не тема: строка только со словом срочности уходит в «разобрать вручную», но срочная.
        s = dict(m.load_slovar(), **{m.URGENT: [(w, m.roots(w)) for w in urgent]})
        self.assertEqual(m.classify("Пахнет газом у входа", s)[0], m.MANUAL)
        self.assertEqual(m.urgent_words("Пахнет газом у входа", s), ["пахнет"])
        self.assertEqual(m.urgent_words("Колонка сломана", s), [])
        self.assertEqual(m.urgent_words("Колонка жанында бензин иісі шығады", s), ["иіс"])
        self.assertEqual(m.urgent_words("Пахнет бензином, похоже на утечку", s), ["пахнет", "утечка"])

    def test_urgent_rows_case_week1(self):
        """Неделя 1 кейса: строка 48 (запах бензина и утечка на АЗС-05) срочная, красная, первая на карточке.
        Лишняя срочная одна: 13 «запах на всю площадку» от мусорных баков (слово «запах» без топлива).
        Срочное лучше проверить лишний раз, чем пропустить утечку."""
        self.with_urgent()
        r = self.run_case("obrashcheniya.xlsx")
        hot = r["sverka"]["срочно"]
        self.assertEqual(hot[0], "48")
        self.assertLessEqual(len(hot) - 1, 2, hot)
        self.assertEqual(set(hot), {"48", "13"})
        self.assertEqual(r["status"], "red")
        self.assertTrue(r["summary"].startswith("Срочно: 2 (№ 48, 13). Обращений 80"), r["summary"])
        self.assertEqual(r["rows"][0][0], "48")
        self.assertEqual(r["rows"][0][2], "оборудование")  # тема своя: «колонка», срочность не тема
        self.assertEqual(r["row_status"][:2], ["red", "red"])
        self.assertEqual(r["row_note"][0], "срочно: пахнет, утечка")
        self.assertEqual(r["row_note"][1], "срочно: запах; разобрать вручную")  # мусор: слов темы нет
        self.assertNotIn("red", r["row_status"][2:])
        self.assertEqual(r["sverka"]["по строкам"]["48"], "оборудование")
        urgent_sheet = sheet(self.c.out / "obrashcheniya-po-temam.xlsx", "Срочно")
        self.assertEqual(urgent_sheet[0], ["№", "АЗС", "Тема", "Текст", "Слово срочности"])
        self.assertEqual([row[0] for row in urgent_sheet[1:]], ["48", "13"])
        self.assertEqual(urgent_sheet[1][4], "пахнет, утечка")
        table = {row[0]: row for row in sheet(self.c.out / "obrashcheniya-po-temam.xlsx", "Обращения")[1:]}
        self.assertEqual(table["48"][4:], ["колонка, колонок", "срочно: пахнет, утечка"])
        self.assertEqual(table["13"][4:], [None, "срочно: запах; нет слов из словаря"])

    def test_urgent_rows_case_week2(self):
        """Неделя 2 кейса, словарь недели 1: строка 34 (шипит и пахнет газом на АЗС-09) срочная и первая.
        Лишняя срочная одна: 17 «машина дымит» после дизеля (слово «дым»: это качество топлива, не пожар)."""
        self.with_urgent()
        r = self.run_case("obrashcheniya_nedelya2.xlsx")
        hot = r["sverka"]["срочно"]
        self.assertEqual(hot[0], "34")
        self.assertLessEqual(len(hot) - 1, 2, hot)
        self.assertEqual(set(hot), {"34", "17"})
        self.assertEqual(r["rows"][0][0], "34")
        self.assertEqual(r["row_status"][0], "red")
        self.assertTrue(r["row_note"][0].startswith("срочно: пахнет"), r["row_note"][0])
        self.assertTrue(r["summary"].startswith("Срочно: 2 (№ 34, 17)."), r["summary"])

    def test_urgent_word_from_skill_and_key_spelling(self):
        """Слово срочности из скилла человека дописывается в тот же раздел; имя раздела без учёта регистра."""
        m = self.c.module()
        data = dict(SLOVAR, **{"Срочно (не тема)": ["курил"]})
        (self.c.proc / "slovar.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        s = m.load_slovar()
        self.assertIn(m.URGENT, s)
        self.assertNotIn("Срочно (не тема)", s)
        self.assertEqual(m.urgent_words("Охранник курил прямо у колонок", s), ["курил"])
        self.assertEqual(m.classify("Охранник курил прямо у колонок", s)[0], "оборудование")


@unittest.skipUnless(HAVE_LIBS, "нет openpyxl или python-docx")
class Word(unittest.TestCase):
    def setUp(self):
        self.c = Cabinet(compose("word", "main_word"))
        self.addCleanup(self.c.close)

    def texts(self, path):
        import docx
        d = docx.Document(path)
        m = self.c.module()
        return [p.text for p in m.paragraphs(d)], d

    def make_template(self, with_table=True):
        """Шаблон, где Word разбил поле на куски, поле в таблице и в колонтитуле, поле-таблица."""
        import docx
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run("Кому: {{Наименование ")
        r = p.add_run("поставщика}}")
        r.bold = True
        d.add_paragraph("Сумма: {{Сумма}} тенге, дата {{дата}}.")
        if with_table:
            d.add_paragraph("{{Таблица}}")
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "Подпись: {{Фамилия И. О.}}"
        d.sections[0].header.is_linked_to_previous = False
        d.sections[0].header.paragraphs[0].text = "Исх. от {{дата}}"
        d.save(self.c.proc / "shablon.docx")

    def test_fill_fields_split_runs_tables_header(self):
        self.make_template()
        m = self.c.module()
        path, missing = m.fill_docx({"Наименование поставщика": "ТОО «Ромашка»", "дата": "21.09.2026",
                                     "Сумма": m.money(1250000), "Таблица": [["Документ", "Сумма"], ["№ 41", "6 195 000,00"]]},
                                    "pismo-01.docx")
        self.assertEqual(path, "output/proba/pismo-01.docx")
        self.assertEqual(missing, ["Фамилия И. О."])
        texts, d = self.texts(self.c.out / "pismo-01.docx")
        self.assertIn("Кому: ТОО «Ромашка»", texts)
        self.assertIn("Сумма: 1 250 000,00 тенге, дата 21.09.2026.", texts)
        self.assertIn("Подпись: [нет данных]", texts)
        self.assertIn("Исх. от 21.09.2026", texts)
        self.assertNotIn("{{Таблица}}", texts)
        self.assertEqual([c.text for c in d.tables[0].rows[1].cells], ["№ 41", "6 195 000,00"])
        self.assertFalse(any("{{" in t for t in texts))

    def test_first_page_and_even_headers_content_control_and_formatting(self):
        """Бланк с отдельным колонтитулом первой страницы, поле в элементе управления содержимым,
        поле, разбитое на куски с разным оформлением: всё заполнено, жирный и курсив вне поля на месте."""
        import docx
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        d = docx.Document()
        p = d.add_paragraph()
        a = p.add_run("Жирное начало ")
        a.bold = True
        p.add_run("и поле {{Наименование ")
        b = p.add_run("поставщика}}")
        b.italic = True
        c = p.add_run(" и курсивный хвост")
        c.italic = True
        s = d.sections[0]
        s.different_first_page_header_footer = True
        s.first_page_header.paragraphs[0].text = "Бланк, исх. от {{дата}}"
        d.settings.odd_and_even_pages_header_footer = True
        s.even_page_footer.paragraphs[0].text = "Чётная страница {{дата}}"
        d.element.body.insert(0, parse_xml(
            f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p><w:r><w:t>В элементе: {{{{дата}}}}</w:t></w:r></w:p></w:sdtContent></w:sdt>'))
        d.save(self.c.proc / "shablon.docx")
        m = self.c.module()
        _, missing = m.fill_docx({"Наименование поставщика": "ТОО «Ромашка»", "дата": "21.09.2026"}, "pismo-01.docx")
        self.assertEqual(missing, [])
        out = docx.Document(self.c.out / "pismo-01.docx")
        sec = out.sections[0]
        self.assertEqual(sec.first_page_header.paragraphs[0].text, "Бланк, исх. от 21.09.2026")
        self.assertEqual(sec.even_page_footer.paragraphs[0].text, "Чётная страница 21.09.2026")
        texts, _ = self.texts(self.c.out / "pismo-01.docx")
        self.assertIn("В элементе: 21.09.2026", texts)
        runs = [(r.text, r.bold, r.italic) for r in out.paragraphs[0].runs if r.text]
        self.assertEqual(runs[0], ("Жирное начало ", True, None))
        self.assertEqual(runs[1], ("и поле ТОО «Ромашка»", None, None))
        self.assertEqual(runs[-1], (" и курсивный хвост", None, True))

    def test_field_out_of_reach_goes_to_missing(self):
        """Поле там, куда замена не достаёт (перемещённый текст в режиме правок): документ человеку."""
        import docx
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        d = docx.Document()
        d.add_paragraph("Обычное {{дата}}")
        d.element.body.insert(0, parse_xml(
            f'<w:p {nsdecls("w")}><w:moveTo w:id="1" w:author="a"><w:r><w:t>Перенесено {{{{Подпись}}}}</w:t></w:r></w:moveTo></w:p>'))
        d.save(self.c.proc / "shablon.docx")
        m = self.c.module()
        _, missing = m.fill_docx({"дата": "21.09.2026"}, "pismo-01.docx")
        self.assertEqual(missing, ["Подпись"])

    def test_table_field_inside_cell_keeps_paragraph_after(self):
        import docx
        d = docx.Document()
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "{{Таблица}}"
        d.save(self.c.proc / "shablon.docx")
        m = self.c.module()
        m.fill_docx({"Таблица": [["Документ", "Сумма"], ["№ 41", "1"]]}, "pismo-01.docx")
        out = docx.Document(self.c.out / "pismo-01.docx")
        tc = out.tables[0].cell(0, 0)._tc
        kids = [k.tag.split("}")[1] for k in tc if not k.tag.endswith("}tcPr")]
        self.assertEqual(kids, ["tbl", "p"])  # ячейка кончается абзацем, иначе Word считает файл испорченным
        self.assertEqual(out.tables[0].cell(0, 0).tables[0].cell(1, 0).text, "№ 41")

    def test_template_from_person_file_brackets_become_fields(self):
        shutil.copy(CASE / "template_pismo_postavshchiku.docx", self.c.input / "template.docx")
        code = (self.c.proc / "run.py").read_text(encoding="utf-8").replace('SOURCE_TEMPLATE = ""', 'SOURCE_TEMPLATE = "template.docx"')
        (self.c.proc / "run.py").write_text(code, encoding="utf-8")
        m = self.c.module()
        m.ensure_template()
        texts, _ = self.texts(self.c.proc / "shablon.docx")
        joined = "\n".join(texts)
        self.assertIn("{{Наименование поставщика}}", joined)
        self.assertIn("Исх. № ____ от {{дата}}", joined)
        self.assertNotIn("[", joined)
        before = (self.c.proc / "shablon.docx").read_bytes()
        m.ensure_template()  # готовый шаблон не переделывается: человек мог поправить его в Word
        self.assertEqual((self.c.proc / "shablon.docx").read_bytes(), before)

    def test_template_from_skill_text(self):
        m = self.c.module()
        m.ensure_template()
        texts, _ = self.texts(self.c.proc / "shablon.docx")
        self.assertEqual(texts, m.SHABLON_TEXT)

    def write_rows(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Поставщик", "Сумма"])
        ws.append(["ТОО «Ромашка»", 1250000])
        ws.append(["ТОО «Василёк»", None])
        wb.save(self.c.input / "rashozhdeniya.xlsx")

    def edit_program(self, old, new):
        code = (self.c.proc / "run.py").read_text(encoding="utf-8")
        self.assertEqual(code.count(old), 1, f"в образце нет ровно одной строки {old!r}")
        (self.c.proc / "run.py").write_text(code.replace(old, new), encoding="utf-8")

    def type_into_template(self, old, new):
        """Человек открыл shablon.docx в Word и напечатал своё вместо поля."""
        import docx
        d = docx.Document(self.c.proc / "shablon.docx")
        m = self.c.module()
        found = [p for p in m.paragraphs(d) if old in p.text]
        self.assertTrue(found, f"в шаблоне нет {old}")
        for p in found:
            p.runs[0].text = p.text.replace(old, new)
            for r in p.runs[1:]:
                r.text = ""
        d.save(self.c.proc / "shablon.docx")

    def test_example_main_one_document_per_row(self):
        self.write_rows()
        self.make_template(with_table=False)
        proc = self.c.run()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        r = self.c.result()
        # Подписи в программе нет: пока поле в шаблоне, каждое письмо человеку.
        self.assertEqual(r["row_status"], ["info", "info"])
        texts, _ = self.texts(self.c.out / "pismo-01.docx")
        self.assertIn("Подпись: [нет данных]", texts)
        rows = manual_rows(self.c.out / "razobrat-vruchnuyu.xlsx")
        self.assertEqual([x[-1] for x in rows[1:]], ["поля с пометкой: Фамилия И. О.", "поля с пометкой: Сумма, Фамилия И. О."])
        # Человек один раз напечатал подпись в шаблоне вместо поля: дальше она обычный текст.
        self.type_into_template("{{Фамилия И. О.}}", "А. Иванов")
        proc = self.c.run()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        r = self.c.result()
        self.assertEqual(r["row_status"], ["info", None])
        self.assertEqual(r["rows"][0][0], "ТОО «Василёк»")
        self.assertIn("Разобрать вручную: 1", r["summary"])
        self.assertEqual(r["sverka"]["документов"], 2)
        self.assertEqual(r["sverka"]["по строкам"], {"ТОО «Василёк»": "разобрать вручную", "ТОО «Ромашка»": "1 250 000,00"})
        self.assertTrue((self.c.out / "pismo-01.docx").is_file() and (self.c.out / "pismo-02.docx").is_file())
        texts, _ = self.texts(self.c.out / "pismo-01.docx")
        self.assertIn("Подпись: А. Иванов", texts)
        texts, _ = self.texts(self.c.out / "pismo-02.docx")
        self.assertIn("Сумма: [нет данных] тенге, дата 21.09.2026.", texts)
        rows = manual_rows(self.c.out / "razobrat-vruchnuyu.xlsx")
        self.assertEqual(rows[1][-1], "поля с пометкой: Сумма")
        if NODE:
            code, out = self.c.node("list")
            self.assertEqual(code, 0, out)
            self.assertNotIn("Ошибка", out)
            code, out = self.c.node("compare", "proba")  # эталона нет, но проверка программы прошла
            self.assertEqual(code, 2, out)
            self.assertNotIn("нельзя", out)

    def test_skill_marker_field_keeps_document_manual(self):
        """Поле, которое скилл помечает «[уточнить: ...]», отдаёт письмо человеку и при вписанной подписи."""
        self.write_rows()
        self.make_template(with_table=False)
        self.c.module().ensure_template()
        self.type_into_template("{{Фамилия И. О.}}", "А. Иванов")
        import docx
        d = docx.Document(self.c.proc / "shablon.docx")
        d.add_paragraph("Ответить до: {{Срок ответа}}")
        d.save(self.c.proc / "shablon.docx")
        proc = self.c.run()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        r = self.c.result()
        self.assertEqual(r["row_status"], ["info", "info"])
        texts, _ = self.texts(self.c.out / "pismo-01.docx")
        self.assertIn("Ответить до: [уточнить: срок ответа]", texts)
        rows = manual_rows(self.c.out / "razobrat-vruchnuyu.xlsx")
        self.assertEqual([x[-1] for x in rows[1:]], ["поля с пометкой: Срок ответа", "поля с пометкой: Срок ответа, Сумма"])

    def test_marker_from_skill_documents_go_to_manual(self):
        """Скилл пишет «[уточнить]»: та же пометка у полей без данных, письма с ней человеку."""
        self.write_rows()
        self.make_template(with_table=False)
        self.edit_program('NO_DATA = "[нет данных]"', 'NO_DATA = "[уточнить]"')
        proc = self.c.run()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        r = self.c.result()
        self.assertEqual(r["row_status"], ["info", "info"])
        self.assertIn("Разобрать вручную: 2", r["summary"])
        self.assertEqual(r["sverka"]["по строкам"], {"ТОО «Ромашка»": "разобрать вручную", "ТОО «Василёк»": "разобрать вручную"})
        texts, _ = self.texts(self.c.out / "pismo-01.docx")
        self.assertIn("Подпись: [уточнить]", texts)
        texts, _ = self.texts(self.c.out / "pismo-02.docx")
        self.assertIn("Сумма: [уточнить] тенге, дата 21.09.2026.", texts)
        rows = manual_rows(self.c.out / "razobrat-vruchnuyu.xlsx")
        self.assertEqual([x[-1] for x in rows[1:]], ["поля с пометкой: Фамилия И. О.", "поля с пометкой: Сумма, Фамилия И. О."])

    def test_marker_in_value_or_template_text(self):
        """Пометка с пояснением в значении поля и пометка в самом тексте шаблона: документ тоже человеку."""
        self.make_template(with_table=False)
        m = self.c.module()
        values = {"Наименование поставщика": "ТОО «Ромашка»", "дата": "21.09.2026", "Сумма": "1 250 000,00",
                  "Фамилия И. О.": "[уточнить: подписант]"}
        _, missing = m.fill_docx(values, "pismo-01.docx")
        self.assertEqual(missing, ["Фамилия И. О."])
        texts, _ = self.texts(self.c.out / "pismo-01.docx")
        self.assertIn("Подпись: [уточнить: подписант]", texts)
        _, missing = m.fill_docx(dict(values, **{"Фамилия И. О.": "А. Иванов [1]"}), "pismo-02.docx")
        self.assertEqual(missing, [])  # скобки без пометки: не пометка
        import docx
        d = docx.Document(self.c.proc / "shablon.docx")
        d.add_paragraph("Срок оплаты: [уточнить: срок оплаты]. Договор: [нет данных]")
        d.save(self.c.proc / "shablon.docx")
        _, missing = m.fill_docx(dict(values, **{"Фамилия И. О.": "А. Иванов"}), "pismo-03.docx")
        self.assertEqual(missing, ["[нет данных]", "[уточнить: срок оплаты]"])

    @unittest.skipUnless(NODE, "нет node")
    def test_samples_pass_program_check(self):
        """Образцы справки не содержат того, что сверка запрещает программе процесса."""
        for parts in ((), ("slovar", "main_slovar"), ("word", "main_word"), ("slovar", "word", "main_word"),
                      ("word", "dva", "main_dva")):
            c = Cabinet(compose(*parts))
            self.addCleanup(c.close)
            code, out = c.node("compare", "proba")
            self.assertNotIn("нельзя", out, parts)


def sheet(path, name):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return [list(r) for r in wb[name].iter_rows(values_only=True)]
    finally:
        wb.close()


# Отличия договора поставщика от типовой формы (задача 9 кейса): то, что программа должна найти сама.
CASE9_DIFFS = {
    "4.2": "изменён",
    "нет/4.4": "добавлен",
    "5.2": "изменён (мелкая правка)",
    "6.1": "изменён",
    "6.2": "изменён",
    "9/нет": "удалён",
}
FORM_FILE, DOC_FILE = "dogovor_tipovaya_forma.docx", "dogovor_postavshchika.docx"
RISK = "оценить риск: решает юрист"
HINT = "Выход: положите в папку процесса русский текст договора отдельным документом Word с нумерованными разделами."
KZ = ["Тараптар осы шарттың талаптарын орындауға міндетті.", "Жеткізуші тауарды уақытында жеткізеді.",
      "Сатып алушы тауар құнын төлейді.", "Дауларды тараптар келіссөз арқылы шешеді."]


def bilingual(russian_left=False, one_row=False, other=None, mixed=False):
    """Текст до первой таблицы (спецификации) в таблицу в два столбца: русский абзац и рядом казахский
    (other(текст): свой текст второго столбца; mixed: казахская фраза и в русском столбце).
    По строке на абзац или всё в одной строке."""
    def fn(d):
        from docx.oxml.ns import qn
        body = d.element.body
        first = next(e for e in body if e.tag == qn("w:tbl"))
        paras = []
        for e in body:
            if e is first:
                break
            if e.tag == qn("w:p"):
                paras.append(e)
        table = d.add_table(rows=1 if one_row else len(paras), cols=2)
        for k, p in enumerate(paras):
            cells = table.rows[0 if one_row else k].cells
            ru, kz = (cells[0], cells[1]) if russian_left else (cells[1], cells[0])
            text = "".join(x.text or "" for x in p.iter(qn("w:t")))
            ru._tc.append(p)
            if mixed:
                ru.add_paragraph(KZ[k % len(KZ)])
            kz.add_paragraph(other(text) if other else KZ[k % len(KZ)])
        first.addprevious(table._tbl)
    return fn


def set_text(p, text):
    """Текст абзаца одним куском, с оформлением первого куска."""
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(text)


def case_variant(name, dst, change=None):
    """Копия файла кейса 9, изменённая функцией change(документ)."""
    import docx
    d = docx.Document(str(CASE / name))
    if change:
        change(d)
    d.save(str(dst))
    return dst


def headings(change):
    """Правка заголовков разделов (стиль Heading 2): change(абзац, номер, название)."""
    def fn(d):
        for p in d.paragraphs:
            m = re.match(r"(\d+)\.\s+(.*)", p.text)
            if p.style.name == "Heading 2" and m:
                change(p, int(m.group(1)), m.group(2))
    return fn


@unittest.skipUnless(HAVE_LIBS, "нет openpyxl или python-docx")
class TwoDocuments(unittest.TestCase):
    def setUp(self):
        self.c = Cabinet(compose("word", "dva", "main_dva"))
        self.addCleanup(self.c.close)

    def put_case(self, *names):
        for name in names or (FORM_FILE, DOC_FILE):
            shutil.copy(CASE / name, self.c.input / name)

    def run_variant(self, change_form=None, change_doc=None):
        """Программа на копиях файлов кейса с правками; result.json."""
        case_variant(FORM_FILE, self.c.input / FORM_FILE, change_form)
        case_variant(DOC_FILE, self.c.input / DOC_FILE, change_doc)
        proc = self.c.run()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        return self.c.result()

    def spec_of(self, change_doc):
        """Проверка спецификации договора кейса с правкой."""
        m = self.c.module()
        doc = m.read_document(case_variant(DOC_FILE, self.c.input / DOC_FILE, change_doc))
        return m.check_spec(doc, dict(m.contract_card(doc)).get(m.SUM_CARD))

    def test_rule_blanks_numbers_minor_edit(self):
        m = self.c.module()
        self.assertTrue(m.same_text("г. Астана, [дата]", "г. Астана, 18.09.2026"))
        self.assertTrue(m.same_text("Исх. № ____ от «___» ________ 20__ г.", "Исх. № 41 от «18» сентября 2026 г."))
        self.assertTrue(m.same_text("Сумма «  » тенге", "Сумма «120 000» тенге"))
        for blank in ("<срок>", "{срок}", "……", "......", "«...»"):
            self.assertTrue(m.same_text(f"Срок поставки: {blank} дней.", "Срок поставки: 5 дней."), blank)
        self.assertTrue(m.same_text("Цена «ТОВАРА» фиксируется", "цена товара  фиксируется."))  # регистр, кавычки, пробелы
        self.assertFalse(m.same_text("г. Астана, [дата]", "г. Алматы, 18.09.2026"))
        self.assertEqual(m.numbers_in("составляет 46\u00a0400\u00a0000,00 тенге"), [46400000.0])
        self.assertEqual(m.numbers_in("в течение 3 (трёх) рабочих дней"), m.numbers_in("в течение трёх рабочих дней"))
        self.assertEqual(m.numbers_in("120 200 000 (сто двадцать миллионов двести тысяч) тенге"), [120200000.0])
        _, st = m.pair_status("уведомляет в течение 3 (трёх) рабочих дней.", "уведомляет в течение трёх рабочих дней.")
        self.assertEqual(st, "изменён (мелкая правка)")
        # Числа, отрицание, «более/менее», сторона: смысл другой, это не мелкая правка.
        for before, after in (
            ("неустойку 0,1 % в день, но не более 10 % от стоимости партии.",
             "неустойку 0,1 % в день, но не более 5 % от стоимости партии."),
            ("неустойку 0,1 % в день, но не более 10 % от стоимости непоставленного Товара по договору.",
             "неустойку 0,1 % в день, но не менее 10 % от стоимости непоставленного Товара по договору."),
            ("За просрочку поставки Поставщик уплачивает неустойку 0,1 % от стоимости непоставленного Товара за каждый день",
             "За просрочку поставки Покупатель уплачивает неустойку 0,1 % от стоимости непоставленного Товара за каждый день"),
            ("Поставщик вправе расторгнуть договор в одностороннем порядке, уведомив другую сторону за 30 дней.",
             "Поставщик не вправе расторгнуть договор в одностороннем порядке, уведомив другую сторону за 30 дней."),
        ):
            score, st = m.pair_status(before, after)
            self.assertGreaterEqual(score, m.MINOR, after)
            self.assertEqual(st, "изменён", after)

    def test_field_filled_with_clause_goes_to_person(self):
        m = self.c.module()
        suspicious = (
            ("Поставщик: [наименование, БИН, адрес, банковские реквизиты].",
             "Поставщик: ТОО «X». Стороны договорились, что раздел 6 не применяется."),
            ("3.2. Срок поставки партии: [срок] с даты заявки.",
             "3.2. Срок поставки партии: 5 календарных дней, при этом Поставщик вправе в одностороннем порядке "
             "продлить срок поставки на 30 дней с даты заявки."),
            ("4.1. Общая сумма договора составляет [сумма] тенге с учётом НДС и определяется Спецификацией.",
             "4.1. Общая сумма договора составляет 120 200 000 тенге без учёта НДС; НДС уплачивается Покупателем "
             "сверх суммы, сумма в тенге с учётом НДС и определяется Спецификацией."),
        )
        for form, doc in suspicious:
            self.assertTrue(m.same_text(form, doc), doc)  # текст формы на месте, отличием не считается
            problems = m.field_problems(form, doc)
            self.assertEqual(len(problems), 1, doc)
            self.assertIn("нет ли в нём новых условий", problems[0])
        fine = (
            ("г. Астана, [дата]", "г. Астана, 18.09.2026"),
            ("в лице [должность], действующего", "в лице директора Иванова А. Б., действующего"),
            ("Покупатель: ТОО «Дала Ойл», [адрес].", "Покупатель: ТОО «Дала Ойл», г. Астана, ул. Абая, 10."),
            ("составляет [сумма] тенге", "составляет 120 200 000 (сто двадцать миллионов двести тысяч) тенге"),
        )
        for form, doc in fine:
            self.assertEqual(m.field_problems(form, doc), [], doc)
        self.assertEqual(m.field_problems("г. Астана, [дата]", "г. Астана, [дата]"), ["поле [дата] не заполнено"])

    def test_spec_numbers_not_read_go_to_person(self):
        m = self.c.module()
        head = ["№", "Товар", "Ед.", "Количество", "Цена за ед.", "Сумма"]
        doc = [
            {"heading": "", "items": [], "tables": []},
            {"heading": "Приложение 1. Спецификация", "items": [], "tables": [[
                head, ["1", "Бензин АИ-92", "л", "200 000", "228,00", "45 600 000,00"],
                ["2", "Бензин АИ-95", "л", "[количество]", "279,00", "н/д"]]]},
            {"heading": "Приложение 2. Техническая спецификация", "items": [], "tables": [[
                ["Товар", "Количество"], ["Бензин АИ-92", "200 000"]]]},
        ]
        res = m.check_spec(doc, "4.1. Общая сумма договора составляет сто двадцать миллионов тенге.")
        unread = dict(res["unread"])
        self.assertEqual(unread["спецификация, Бензин АИ-95"], "не число: количество «[количество]», сумма «н/д»")
        self.assertEqual(unread["Сумма договора, п. 4.1"], "в тексте пункта нет суммы числом")  # прописью не считается
        self.assertIn("нет строки «Итого» или «Всего»", unread["спецификация"])
        self.assertEqual(res["rows"][0], ["Бензин АИ-92", 200000, 228, 45600000, 45600000, 0])
        self.assertEqual(res["rows"][-2], ["Итого по пересчёту", None, None, None, None, None])  # пропуск: не пересчитать
        self.assertIsNone(res["diff"])
        self.assertEqual(res["tech"], [["Бензин АИ-92", 200000, 200000, 0], ["Бензин АИ-95", None, None, None]])

    def test_spec_totals_vsego_merged_missing_vat(self):
        def vsego(d):  # «Всего» в столбце №, ячейка товара пустая
            row = d.tables[0].rows[-1]
            set_text(row.cells[0].paragraphs[0], "Всего")
            set_text(row.cells[1].paragraphs[0], "")
        res = self.spec_of(vsego)
        self.assertIn(["Всего", None, None, 120200000, 118200000, 2000000], res["rows"])
        self.assertEqual(res["diff"], 2000000)

        def merged(d):  # «Итого» в ячейке, объединённой на пять столбцов
            row = d.tables[0].rows[-1]
            cell = row.cells[0].merge(row.cells[4])
            for p in cell.paragraphs[1:]:
                p._p.getparent().remove(p._p)
            set_text(cell.paragraphs[0], "Итого")
        res = self.spec_of(merged)
        self.assertIn(["Итого", None, None, 120200000, 118200000, 2000000], res["rows"])
        self.assertEqual(res["unread"], [])

        def no_total(d):
            tr = d.tables[0].rows[-1]._tr
            tr.getparent().remove(tr)
        res = self.spec_of(no_total)
        self.assertIsNone(res["diff"])
        self.assertIn("нет строки «Итого» или «Всего»", dict(res["unread"])["спецификация"])
        self.assertIn(["Сумма договора, п. 4.1", None, None, 120200000, 118200000, 2000000], res["rows"])

        def vat(d):  # «Итого без НДС», «НДС 12 %», «Итого с НДС»
            t = d.tables[0]
            last = t.rows[-1]._tr
            for _ in range(2):
                last.addnext(copy.deepcopy(last))
            for row, (name, value) in zip(t.rows[-3:], (("Итого без НДС", "107 321 428,57"), ("НДС 12 %", "12 878 571,43"),
                                                        ("Итого с НДС", "120 200 000,00"))):
                set_text(row.cells[1].paragraphs[0], name)
                set_text(row.cells[5].paragraphs[0], value)
        res = self.spec_of(vat)
        self.assertIsNone(res["diff"])
        note = dict(res["unread"])["спецификация"]
        self.assertIn("несколько итоговых строк (Итого без НДС, НДС 12 %, Итого с НДС)", note)
        self.assertIn("118 200 000,00", note)
        self.assertEqual([p for p in res["problems"] if "НДС" in p], [])  # строки НДС товаром не считаются
        self.assertIn("Дизельное топливо: 46 400 000,00 при пересчёте 44 400 000,00", res["problems"])

    def test_headings_auto_numbered_roman_and_final_dot(self):
        """Номер раздела списком Word, римский номер, «4. Название.» с точкой без стиля: тот же разбор."""
        romans = "I II III IV V VI VII VIII IX X XI".split()

        def dot(p, n, title):
            p.style = p.part.document.styles["Normal"]
            set_text(p, f"{n}. {title}.")
        variants = {
            "нумерация списком": (headings(lambda p, n, title: set_text(p, title)), "Конфиденциальность/нет"),
            "римские номера": (headings(lambda p, n, title: set_text(p, f"{romans[n - 1]}. {title}")), "IX/нет"),
            "точка в конце": (headings(dot), "9/нет"),
        }
        for name, (change, deleted) in variants.items():
            with self.subTest(name):
                r = self.run_variant(change, change)
                expected = {k: v for k, v in CASE9_DIFFS.items() if k != "9/нет"} | {deleted: "удалён"}
                self.assertEqual(r["sverka"]["по строкам"], expected)
                self.assertEqual(r["sverka"]["итоги"]["удалено"], 1)
                status = {row[0]: row[1] for row in r["rows"]}
                self.assertEqual(status["10.1/9.1"], "совпадает")
                self.assertFalse([row for row in r["rows"] if str(row[0]).startswith(("документ:", "карточка:"))])
                self.assertIn("Сумма договора, п. 4.1: 120 200 000,00 при пересчёте 118 200 000,00", r["summary"])

    def test_sections_not_found_go_to_person(self):
        def plain_headings(p, n, title):  # заголовок обычным жирным текстом, без номера и стиля
            p.style = p.part.document.styles["Normal"]
            set_text(p, title.upper())
            p.runs[0].bold = True
        r = self.run_variant(headings(plain_headings), headings(plain_headings))
        self.assertEqual(r["status"], "yellow")
        self.assertIn("Сверка неполная", r["summary"])
        why = manual_rows(self.c.out / "razobrat-vruchnuyu.xlsx")
        self.assertTrue(any(row[0] == f"документ: {DOC_FILE}" and "разделы не распознаны" in row[-1] for row in why))

    def test_contract_in_table_goes_to_person(self):
        """Договор свёрстан таблицей, второго столбца с текстом нет: вне таблиц текста нет, подсказка, что делать."""
        def in_table(d):
            from docx.oxml.ns import qn
            body = d.element.body
            paras = [e for e in body if e.tag == qn("w:p")]
            table = d.add_table(rows=1, cols=2)
            left, right = table.rows[0].cells
            for p in paras:
                right._tc.append(p)
            left.paragraphs[0].text = "Қазақ тіліндегі мәтін"
            body.insert(0, table._tbl)
        r = self.run_variant(in_table, in_table)
        self.assertEqual(r["status"], "yellow")
        self.assertIn("Сверка неполная", r["summary"])
        self.assertEqual(r["sverka"]["по строкам"], {})
        why = manual_rows(self.c.out / "razobrat-vruchnuyu.xlsx")
        row = next(row for row in why if row[0] == f"документ: {DOC_FILE}")
        self.assertIn("вне таблиц нет текста", row[-1])
        self.assertTrue(row[-1].endswith(HINT), row[-1])

    def test_bilingual_two_columns_read_by_russian_column(self):
        """Договор на двух языках в два столбца: программа читает русский столбец, отличия те же, что у кейса 9."""
        variants = {
            "оба двуязычные": (bilingual(), bilingual(russian_left=True, one_row=True)),
            "только договор": (None, bilingual()),
            "только форма": (bilingual(one_row=True), None),
        }
        for name, (change_form, change_doc) in variants.items():
            with self.subTest(name):
                r = self.run_variant(change_form, change_doc)
                self.assertEqual(r["sverka"]["по строкам"], CASE9_DIFFS)
                self.assertEqual(r["sverka"]["итоги"]["разница спецификации"], 2000000)
                self.assertNotIn("Сверка неполная", r["summary"])
                phrase = "Договор двуязычный: сверялся русский текст." if change_doc else \
                    "Типовая форма двуязычная: сверялся русский текст."
                self.assertIn(phrase, r["summary"])
                manual = manual_rows(self.c.out / "razobrat-vruchnuyu.xlsx")[1:]
                self.assertEqual(len(manual), len(CASE9_DIFFS) + 2, manual)  # как у кейса 9: отличия и два поля
                card = dict(sheet(self.c.out / "sverka-dogovora.xlsx", "Карточка условий")[1:])
                self.assertEqual(card["Номер"], "ПС-2026/118")
                self.assertNotIn("[нет данных]", card.values())

    def test_bilingual_language_unclear_goes_to_person(self):
        """Оба столбца на русском (или оба с казахскими буквами): язык не определить, не угадывать."""
        for name, change in (("оба русские", bilingual(other=lambda text: text)), ("оба с казахскими", bilingual(mixed=True))):
            with self.subTest(name):
                r = self.run_variant(None, change)
                self.assertEqual(r["status"], "yellow")
                self.assertNotIn("двуязычный", r["summary"])
                why = manual_rows(self.c.out / "razobrat-vruchnuyu.xlsx")
                row = next(row for row in why if row[0] == f"документ: {DOC_FILE}")
                self.assertIn("вне таблиц нет текста", row[-1])
                self.assertIn("язык столбцов не определить", row[-1])
                self.assertTrue(row[-1].endswith(HINT), row[-1])

    def test_paired_heading_text_change(self):
        """Одноуровневые пункты «4. Срок действия один год.» программа берёт за заголовок раздела.
        Правка в такой строке («один год и более») видна: строка раздела «изменён» и человеку."""
        def contract(last):
            import docx
            d = docx.Document()
            d.add_paragraph("Договор оказания услуг № [номер]", style="Heading 1")
            d.add_paragraph("г. Астана, [дата]")
            for n, title, point in ((1, "Предмет договора", "Исполнитель оказывает услуги по заявкам Заказчика."),
                                    (2, "Цена услуг", "Цена услуг указана в Приложении 1."),
                                    (3, "Порядок оплаты", "Заказчик оплачивает услуги в течение 30 дней.")):
                d.add_paragraph(f"{n}. {title}")
                d.add_paragraph(f"{n}.1. {point}")
            d.add_paragraph(f"4. Срок действия {last}.")
            d.add_paragraph("5. Споры решаются в суде г. Астаны.")
            return d
        contract("один год").save(self.c.input / FORM_FILE)
        doc = contract("один год и более")
        doc.paragraphs[1].text = "г. Астана, 18.09.2026"
        doc.save(self.c.input / DOC_FILE)
        proc = self.c.run()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        r = self.c.result()
        self.assertEqual(r["sverka"]["по строкам"], {"4": "изменён"})
        self.assertEqual(r["sverka"]["итоги"]["изменено"], 1)
        manual = manual_rows(self.c.out / "razobrat-vruchnuyu.xlsx")[1:]
        row = next(row for row in manual if row[0] == "4")
        self.assertEqual(row[1:], ["изменён", "4. Срок действия один год.", "4. Срок действия один год и более.", RISK])
        m = self.c.module()
        self.assertEqual(m.pair_status("Срок действия договора", "СРОК ДЕЙСТВИЯ ДОГОВОРА")[1], "совпадает")

    def test_form_blank_styles_and_underlined_spaces(self):
        """Поля формы в другом виде: [..] заменены на <..>, {..}, ……, точки, подчёркнутые пробелы. Отличия те же."""
        def blanks(write):
            def fn(d):
                for p in d.paragraphs:
                    if "[" in p.text:
                        parts = re.split(r"\[[^\]]*\]", p.text)
                        for r in list(p.runs):
                            r._r.getparent().remove(r._r)
                        for k, part in enumerate(parts):
                            p.add_run(part)
                            if k < len(parts) - 1:
                                write(p)
            return fn

        def underline(p):
            p.add_run("            ").underline = True
        for name, write in (("<..>", lambda p: p.add_run("<указать>")), ("{..}", lambda p: p.add_run("{сумма}")),
                            ("……", lambda p: p.add_run("…………")), ("точки", lambda p: p.add_run("...........")),
                            ("подчёркнутые пробелы", underline)):
            with self.subTest(name):
                r = self.run_variant(blanks(write))
                self.assertEqual(r["sverka"]["по строкам"], CASE9_DIFFS)

    def test_pairs_do_not_cross(self):
        """5.1 удалён, новый 5.3 о приёмке после 5.2: «удалён» и «добавлен», не «изменён» крест-накрест."""
        def change(d):
            p51 = next(p for p in d.paragraphs if p.text.startswith("5.1."))
            p52 = next(p for p in d.paragraphs if p.text.startswith("5.2."))
            p51._p.getparent().remove(p51._p)
            new = copy.deepcopy(p52._p)
            p52._p.addnext(new)
            from docx.text.paragraph import Paragraph
            set_text(Paragraph(new, p52._parent), "5.3. Приёмка производится представителем Покупателя на основании доверенности.")
        r = self.run_variant(None, change)
        rows = r["sverka"]["по строкам"]
        self.assertEqual(rows["5.1/нет"], "удалён")
        self.assertEqual(rows["нет/5.3"], "добавлен")
        self.assertEqual(rows["5.2"], "изменён (мелкая правка)")

    def test_no_form_clear_stop(self):
        self.put_case(DOC_FILE)
        proc = self.c.run()
        self.assertEqual(proc.returncode, 2, proc.stdout + proc.stderr)
        self.assertIn("нет типовой формы", proc.stderr)
        self.assertNotIn("Traceback", proc.stderr)

    def test_case_9_contract_against_form(self):
        self.put_case()
        proc = self.c.run()
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        r = self.c.result()
        status = {row[0]: row[1] for row in r["rows"]}
        self.assertEqual({k: v for k, v in status.items() if v != "совпадает"}, CASE9_DIFFS)
        # После удалённого раздела 9 разделы 10 и 11 формы стоят в договоре под номерами 9 и 10: не отличие.
        self.assertEqual(status["10.1/9.1"], "совпадает")
        self.assertEqual(status["11, абз. 1/10, абз. 1"], "совпадает")
        self.assertEqual(status["11, абз. 2/10, абз. 2"], "совпадает")
        # Заполненные поля формы (номер, дата, поставщик, сумма, реквизиты) не отличие.
        for key in ("шапка, абз. 1", "шапка, абз. 2", "шапка, абз. 3", "4.1"):
            self.assertEqual(status[key], "совпадает", key)
        # Человеку: шесть отличий (риск) и два поля реквизитов, заполненных несколькими фразами.
        n = len(CASE9_DIFFS) + 2
        self.assertEqual(r["row_status"], ["info"] * n + [None] * (len(r["rows"]) - n))
        self.assertEqual(r["row_note"][:n], ["разобрать вручную"] * n)
        manual = manual_rows(self.c.out / "razobrat-vruchnuyu.xlsx")[1:]
        self.assertEqual(sorted(row[0] for row in manual if row[-1] == RISK), sorted(CASE9_DIFFS))
        self.assertEqual([row[0] for row in manual if row[-1] != RISK], ["11, абз. 1/10, абз. 1", "11, абз. 2/10, абз. 2"])
        self.assertTrue(all("нет ли в нём новых условий" in row[-1] for row in manual if row[-1] != RISK))
        self.assertEqual(r["sverka"]["по строкам"], CASE9_DIFFS)
        self.assertEqual(r["sverka"]["итоги"], {"изменено": 3, "добавлено": 1, "удалено": 1, "мелких правок": 1,
                                                "разница спецификации": 2000000})
        self.assertEqual(r["status"], "yellow")  # цифры не сходятся
        self.assertNotIn("Сверка неполная", r["summary"])
        self.assertIn(f"Разобрать вручную: {n} (файл razobrat-vruchnuyu.xlsx)", r["summary"])

        book = self.c.out / "sverka-dogovora.xlsx"
        self.assertIn("output/proba/sverka-dogovora.xlsx", r["files"])
        spec = {row[0]: row[1:] for row in sheet(book, "Спецификация")[1:]}
        self.assertEqual(spec["Дизельное топливо"], [150000, 296, 46400000, 44400000, 2000000])
        self.assertEqual(spec["Бензин АИ-92"][-1], 0)
        self.assertEqual(spec["Итого"][2:], [120200000, 118200000, 2000000])
        self.assertEqual(spec["Сумма договора, п. 4.1"][2:], [120200000, 118200000, 2000000])
        tech = sheet(book, "Техспецификация")
        self.assertEqual([row for row in tech[1:] if row[-1] != 0], [["Бензин АИ-95", 80000, 85000, -5000]])
        self.assertEqual(len(tech), 5)
        card = dict(sheet(book, "Карточка условий")[1:])
        self.assertEqual(card["Поставщик"], "ТОО «НефтеСнаб Сарыарка»")
        self.assertEqual(card["Покупатель"], "ТОО «Дала Ойл»")
        self.assertEqual(card["Номер"], "ПС-2026/118")
        self.assertEqual(card["Дата"], "18.09.2026")
        self.assertIn("31 декабря 2026", card["Срок действия"])
        self.assertIn("г. Астана", card["Споры"])
        self.assertIn("120 200 000", card["Сумма"])
        self.assertIn("предоплату 50 %", card["Оплата"])
        self.assertNotIn("[нет данных]", card.values())

        if not NODE:
            return
        code, out = self.c.node("list")
        self.assertEqual(code, 0, out)
        self.assertNotIn("Ошибка", out)
        etalon = {"по строкам": CASE9_DIFFS, "итоги": {"изменено": 3, "добавлено": 1, "удалено": 1,
                                                       "мелких правок": 1, "разница спецификации": "2 000 000,00"}}
        (self.c.out / "etalon-skilla.json").write_text(json.dumps(etalon, ensure_ascii=False), encoding="utf-8")
        code, out = self.c.node("compare", "proba")
        self.assertEqual(code, 0, out)
        self.assertIn("Итог сверки: программа дала тот же результат, что скилл.", out)
        self.assertNotIn("Внимание", out)


if __name__ == "__main__":
    unittest.main()
