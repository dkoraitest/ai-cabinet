#!/usr/bin/env python3
"""Быстро заглянуть в таблицу (Excel или выгрузка 1С) перед любым расчётом.

    python tools/read_table.py <файл> [--sheet N|имя] [--rows 10] [--json]
    python tools/read_table.py <файл> [--sum СТОЛБЕЦ] [--by СТОЛБЕЦ] [--where СТОЛБЕЦ=ЗНАЧЕНИЕ]

Что печатает: кодировку и разделитель (для CSV/TXT), листы с числом строк (для XLSX),
строку шапки, названия столбцов, число строк данных и первые строки.
С --sum, --by или --where: итоги по всем строкам (число строк и суммы столбцов, всего и по группам
до двух столбцов, с отбором строк). Ими Claude проверяет расчёт программы процесса независимо от неё.
Объединённая многоуровневая шапка (до 4 строк) разворачивается в названия «Родитель / Дочерний».
--sheet принимает номер листа (с 1) или его имя.

Только стандартная библиотека и openpyxl, без pandas. Большие файлы читаются потоком.
Объединённые ячейки в режиме read_only openpyxl не видны, поэтому диапазоны объединений
берутся прямо из XML листа внутри xlsx (потоковое чтение, без второй загрузки книги).
Код выхода 2 и русское сообщение: файл не найден, тип не поддерживается, старый .xls,
файл защищён паролем.
"""

import argparse
import codecs
import csv
import io
import json
import re
import sys
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from xml.etree import ElementTree as ET

HEADER_SCAN_ROWS = 30
CSV_CANDIDATES = [";", ",", "\t", "|"]
CSV_EXT = {".csv", ".txt", ".tsv"}
XLSX_EXT = {".xlsx", ".xlsm"}
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
DELIM_NAMES = {";": "точка с запятой (;)", ",": "запятая (,)", "\t": "табуляция", "|": "вертикальная черта (|)"}


class UserError(Exception):
    """Понятная человеку ошибка: печатается по-русски, код выхода 2."""


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "да" if v else "нет"
    if isinstance(v, float):
        return str(int(v)) if v == int(v) and abs(v) < 1e15 else repr(v)
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0, 0) else v.isoformat(sep=" ")
    if isinstance(v, (date, time)):
        return v.isoformat()
    return str(v).strip()


# ---------- общий разбор строк: шапка, счёт, первые строки ----------

def pick_header(buffer):
    """Индекс строки шапки: первая, где заполнено не меньше половины ячеек."""
    width = 0
    for row in buffer:
        filled = [i for i, c in enumerate(row) if c != ""]
        if filled:
            width = max(width, filled[-1] + 1)
    for i, row in enumerate(buffer):
        n = sum(1 for c in row if c != "")
        if n and n * 2 >= width and (n >= 2 or width == 1):
            return i
    for i, row in enumerate(buffer):
        if any(c != "" for c in row):
            return i
    return None


NUMERIC_RE = re.compile(r"^[-+]?[\d\s.,:/]+%?$")
MAX_LEVELS = 4


def is_labels(row):
    """Строка из подписей: есть непустые ячейки и среди них нет чисел и дат."""
    cells = [c for c in row if c != ""]
    return bool(cells) and not any(NUMERIC_RE.match(c) for c in cells)


def build_columns(buffer, h, merges):
    """Названия столбцов, число строк шапки и индекс её первой строки.

    Шапка может занимать несколько строк (до 4 уровней). Строка ниже входит в шапку, если
    выше начинается объединение по горизонтали, под которым стоят подписи, или объединение
    вниз заканчивается ровно на этой строке, и в ней одни подписи, а не данные.
    Объединение вниз на много строк (за шапку, в данные) шапку не удлиняет.
    Одноячеечная строка выше шапки считается верхним уровнем, только если под ней есть
    объединение меньшего размера (иначе это заголовок отчёта).
    """
    t = h
    while t > 0 and any(buffer[t - 1]):
        up = [m for m in merges if m[0] == t and m[3] > m[1]]
        if any(q[0] == t + 1 and q[3] > q[1] and p[1] <= q[1] and q[3] <= p[3]
               and (q[1], q[3]) != (p[1], p[3]) for p in up for q in merges):
            t -= 1
        else:
            break
    last = h
    while last + 1 < len(buffer) and last - t + 1 < MAX_LEVELS:
        nxt = buffer[last + 1]
        trig_h = any(m[0] == last + 1 and m[3] > m[1] and any(nxt[c] for c in range(m[1] - 1, min(m[3], len(nxt))))
                     for m in merges)
        trig_v = any(t + 1 <= m[0] <= last + 1 and m[2] == last + 2 for m in merges)
        if (trig_h or trig_v) and is_labels(nxt):
            last += 1
        else:
            break
    block = buffer[t:last + 1]
    width = max(max((i + 1 for i, c in enumerate(r) if c != ""), default=0) for r in block)
    for m in merges:
        if t + 1 <= m[0] <= last + 1:
            width = max(width, m[3])
    grid = [(list(r) + [""] * width)[:width] for r in block]
    for r1, c1, r2, c2 in merges:
        if not t + 1 <= r1 <= last + 1:
            continue
        val = grid[r1 - 1 - t][c1 - 1] if c1 - 1 < width else ""
        for r in range(r1, min(r2, last + 1) + 1):
            for c in range(c1 - 1, min(c2, width)):
                if grid[r - 1 - t][c] == "":
                    grid[r - 1 - t][c] = val
    cols = []
    for i in range(width):
        parts = []
        for r in grid:
            if r[i] and (not parts or parts[-1] != r[i]):
                parts.append(r[i])
        cols.append(" / ".join(parts) if parts else f"(без названия {i + 1})")
    while cols and cols[-1].startswith("(без названия"):
        cols.pop()
    return cols, last - t + 1, t


def scan(row_iter, merges, n_rows):
    """Один проход по строкам (списки строк, пустая строка = пустой список/ячейки).

    Объединённые вниз ячейки в строках данных заполняются значением из верхней ячейки.
    """
    starts = {}
    for m in merges:
        if m[2] > m[0]:
            starts.setdefault(m[0], []).append(m)
    active = []  # (последняя строка, столбец1, столбец2, значение)
    buffer, cols, hrows = [], None, 1
    h = None
    nonempty, data_rows, sample = 0, 0, []

    def take(filled):
        nonlocal data_rows
        if any(filled):
            data_rows += 1
            if len(sample) < n_rows:
                sample.append(filled)

    def resolve():
        nonlocal cols, hrows, h, buffer
        h = pick_header([raw for raw, _ in buffer])
        if h is None:
            return
        cols, hrows, h = build_columns([raw for raw, _ in buffer], h, merges)
        for raw, filled in buffer[h + hrows:]:
            if any(raw):
                take(filled)
        buffer = []

    for idx, row in enumerate(row_iter):
        rn = idx + 1
        raw = [cell_text(c) for c in row]
        filled = raw
        active = [a for a in active if a[0] >= rn]
        if active:
            filled = list(raw)
            for last_row, c1, c2, val in active:
                if len(filled) < c2:
                    filled += [""] * (c2 - len(filled))
                for c in range(c1 - 1, c2):
                    if filled[c] == "":
                        filled[c] = val
        for r1, c1, r2, c2 in starts.get(rn, ()):
            active.append((r2, c1, c2, raw[c1 - 1] if c1 - 1 < len(raw) else ""))
        if any(raw):
            nonempty += 1
        if cols is None:
            buffer.append((raw, filled))
            if len(buffer) >= HEADER_SCAN_ROWS:
                resolve()
            continue
        if any(raw):
            take(filled)
    if cols is None:  # файл короче окна поиска шапки
        resolve()
    width = len(cols) if cols else 0
    sample = [(r + [""] * width)[:width] for r in sample]
    return {
        "header_row": (h + 1) if cols is not None else None,
        "header_rows": hrows,
        "columns": cols or [],
        "data_rows": data_rows,
        "total_rows": nonempty,
        "rows": sample,
    }


# ---------- CSV / TXT ----------

def detect_encoding(path):
    """utf-8-sig (есть BOM), utf-8 или cp1251. Проверка потоком, файл целиком не грузится."""
    with path.open("rb") as f:
        head = f.read(3)
        if head == codecs.BOM_UTF8:
            return "utf-8-sig"
        f.seek(0)
        dec = codecs.getincrementaldecoder("utf-8")()
        try:
            while True:
                chunk = f.read(1 << 20)
                dec.decode(chunk, final=not chunk)
                if not chunk:
                    break
        except UnicodeDecodeError:
            return "cp1251"
    return "utf-8"


def detect_delimiter(lines):
    lines = [ln for ln in lines if ln.strip()]
    if not lines:
        return ";"
    try:
        d = csv.Sniffer().sniff("\n".join(lines), delimiters="".join(CSV_CANDIDATES)).delimiter
        multi = sum(1 for ln in lines if d in ln)
        if d in CSV_CANDIDATES and multi * 2 >= len(lines):
            return d
    except csv.Error:
        pass
    # Запасной вариант: символ, который чаще всего встречается и в большинстве строк.
    def score(d):
        counts = [ln.count(d) for ln in lines]
        return (sum(1 for c in counts if c > 0), sum(counts))
    best = max(CSV_CANDIDATES, key=score)
    return best if score(best)[0] else ";"


def read_csv(path, n_rows):
    enc = detect_encoding(path)
    with path.open("r", encoding=enc, errors="replace", newline="") as f:
        sample_lines = []
        for _ in range(60):
            ln = f.readline()
            if not ln:
                break
            sample_lines.append(ln)
    delim = detect_delimiter(sample_lines)
    with path.open("r", encoding=enc, errors="replace", newline="") as f:
        csv.field_size_limit(1 << 24)
        res = scan(csv.reader(f, delimiter=delim), [], n_rows)
    res.update(type="csv", encoding=enc, delimiter=delim, warnings=[])
    return res


# ---------- XLSX ----------

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
MERGE_RE = re.compile(rb'<(?:\w+:)?mergeCell\s[^>]*?ref="([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?"')


def col_num(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def sheet_merges(path, sheet_index):
    """Диапазоны объединённых ячеек листа: (строка1, столбец1, строка2, столбец2), с 1.

    Читает XML листа потоком. Возвращает (диапазоны, предупреждение или None):
    при неудаче диапазонов нет, а причина попадает в вывод.
    """
    try:
        with zipfile.ZipFile(path) as z:
            wb = ET.fromstring(z.read("xl/workbook.xml"))
            rel_id = list(wb.find("m:sheets", NS))[sheet_index].get(f"{{{NS['r']}}}id")
            rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
            target = next(r.get("Target") for r in rels if r.get("Id") == rel_id)
            inner = target.lstrip("/") if target.startswith("/") else "xl/" + target
            out, tail = set(), b""
            with z.open(inner) as f:
                while True:
                    chunk = f.read(1 << 20)
                    data = tail + chunk
                    for m in MERGE_RE.finditer(data):
                        c2 = m.group(3) or m.group(1)
                        r2 = m.group(4) or m.group(2)
                        out.add((int(m.group(2)), col_num(m.group(1).decode()),
                                 int(r2), col_num(c2.decode())))
                    if not chunk:
                        break
                    tail = data[-300:]  # хвост на случай разрыва тега между кусками; повторы уйдут в set
            return sorted(out), None
    except Exception as e:
        return [], f"не удалось прочитать объединённые ячейки ({type(e).__name__}: {e}), шапка может быть разобрана неточно"


def open_book(path):
    import openpyxl
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except zipfile.BadZipFile:
        raise UserError(f"Файл «{path.name}» не похож на xlsx: он повреждён или закрыт паролем.")


def iter_sheet_rows(ws):
    try:
        ws.reset_dimensions()
    except Exception:
        pass
    for row in ws.iter_rows(values_only=True):
        yield list(row)


def resolve_sheet(names, wanted):
    """Лист по имени или номеру с 1. Номер можно дать строкой («2») или числом (2, из программы процесса)."""
    if wanted is None:
        return 0
    listing = "; ".join(f"{i + 1}. {n}" for i, n in enumerate(names))
    if isinstance(wanted, bool) or not isinstance(wanted, (str, int, float)):
        raise UserError(f"Лист нужно указать номером с 1 или именем. Есть листы: {listing}.")
    if isinstance(wanted, (int, float)):
        if wanted != int(wanted):
            raise UserError(f"Номер листа {wanted} не целый. Есть листы: {listing}.")
        number = int(wanted)
        if 1 <= number <= len(names):
            return number - 1
        if str(number) in names:
            return names.index(str(number))
        raise UserError(f"Листа с номером {number} нет: в файле {len(names)}, номера с 1. Есть листы: {listing}.")
    if wanted in names:
        return names.index(wanted)
    text = wanted.strip()
    if text.isdigit() and 1 <= int(text) <= len(names):
        return int(text) - 1
    raise UserError(f"Лист «{wanted}» не найден. Есть листы: {listing}.")


def read_xlsx(path, sheet, n_rows):
    wb = open_book(path)
    try:
        names = list(wb.sheetnames)
        si = resolve_sheet(names, sheet)
        merges, warn = sheet_merges(path, si)
        res = scan(iter_sheet_rows(wb[names[si]]), merges, n_rows)
        sheets = []
        for i, name in enumerate(names):
            if i == si:
                sheets.append({"index": i + 1, "name": name, "rows": res["total_rows"]})
            else:
                cnt = sum(1 for row in iter_sheet_rows(wb[name]) if any(c not in (None, "") for c in row))
                sheets.append({"index": i + 1, "name": name, "rows": cnt})
    finally:
        wb.close()
    res.update(type="xlsx", sheets=sheets, sheet={"index": si + 1, "name": names[si]},
               merged_ranges=len(merges), warnings=[warn] if warn else [])
    return res


# ---------- вход и вывод ----------

def check_file(path):
    if not path.exists() or not path.is_file():
        raise UserError(f"Файл не найден: {path}")
    ext = path.suffix.lower()
    if ext == ".xls":
        raise UserError("Старый формат .xls не читается. Откройте файл в Excel и сохраните как .xlsx, затем повторите.")
    if ext not in CSV_EXT | XLSX_EXT:
        raise UserError(f"Тип файла «{ext or 'без расширения'}» не поддерживается. Нужны .xlsx, .csv или .txt.")
    if ext in XLSX_EXT:
        with path.open("rb") as f:
            magic = f.read(8)
        if magic == OLE_MAGIC:
            raise UserError(f"Файл «{path.name}» защищён паролем (или это старый Excel). "
                            "Снимите пароль в Excel и сохраните как .xlsx.")
    return ext


def analyze(path, sheet=None, n_rows=10):
    ext = check_file(path)
    res = read_xlsx(path, sheet, n_rows) if ext in XLSX_EXT else read_csv(path, n_rows)
    res["file"] = str(path)
    return res


# ---------- итоги по группам: --sum, --by, --where ----------

MAX_GROUPS = 300
LIMIT_TOTALS = 10 ** 6  # строк для итогов: файл больше, вывод об этом скажет
# Дата и, в выгрузках 1С, время после неё: «14.09.2026», «14.09.2026 0:00:00», «2026-09-14 10:30:00».
TIME_TAIL = r"(?:[ T]\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?"
DATE_DMY = re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})" + TIME_TAIL + "$")
DATE_ISO = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})" + TIME_TAIL + "$")
NUMBER_RE = re.compile(r"^[-+]?\d+(?:\.\d+)?$")


def norm_text(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower().replace("ё", "е")


def to_number(text):
    """Число из ячейки: «1 234,5» или «1234.5», пробелы любые (и неразрывный из 1С). Не число: None.

    Только цифры, знак и одна дробная часть: nan, inf, 1_000 и 1e5 числом не считаются.
    """
    t = "".join(str(text).split()).replace(",", ".")
    return float(t) if NUMBER_RE.match(t) else None


def to_date(text):
    t = str(text).strip()
    m = DATE_DMY.match(t)
    if m:
        d, mo, y = (int(g) for g in m.groups())
    else:
        m = DATE_ISO.match(t)
        if not m:
            return None
        y, mo, d = (int(g) for g in m.groups())
    try:
        return date(y, mo, d)
    except ValueError:
        return None


def find_column(columns, wanted):
    """Номер столбца (с 0) по номеру с 1, по точному названию или по части названия, без учёта регистра."""
    w = norm_text(wanted)
    listing = "; ".join(f"{i + 1}. {c}" for i, c in enumerate(columns))
    if w.isdigit():
        if 1 <= int(w) <= len(columns):
            return int(w) - 1
        raise UserError(f"Столбца с номером {w} нет. Столбцы: {listing}.")
    exact = [i for i, c in enumerate(columns) if norm_text(c) == w]
    found = exact or [i for i, c in enumerate(columns) if w and w in norm_text(c)]
    if len(found) == 1:
        return found[0]
    if not found:
        raise UserError(f"Столбец «{wanted}» не найден. Столбцы: {listing}.")
    raise UserError(f"Под «{wanted}» подходят несколько столбцов, укажите номер. Столбцы: {listing}.")


def in_range(convert, low, high):
    def test(text):
        v = convert(text)
        return v is not None and low <= v <= high
    test.convert = convert  # чем читается значение: по нему считаются строки, где оно не читается
    return test


def equals(value):
    wanted = norm_text(value)
    return lambda text: norm_text(text) == wanted


def parse_where(columns, items):
    """--where «Столбец=значение» или «Столбец=от..до» (даты или числа). Один столбец несколько раз: любое из
    значений; разные столбцы: все условия сразу."""
    conds = {}
    for item in items:
        if "=" not in item:
            raise UserError(f"Отбор «{item}» нужно писать так: Столбец=значение или Столбец=от..до.")
        col, value = item.split("=", 1)
        i = find_column(columns, col)
        low, sep, high = value.partition("..")
        if sep:
            if to_date(low) and to_date(high):
                test = in_range(to_date, to_date(low), to_date(high))
            elif to_number(low) is not None and to_number(high) is not None:
                test = in_range(to_number, to_number(low), to_number(high))
            else:
                raise UserError(f"Отбор «{item}»: от..до нужны две даты (ДД.ММ.ГГГГ) или два числа.")
        else:
            test = equals(value)
        conds.setdefault(i, []).append((item, test))
    return conds


def group_totals(r, sums=(), by=(), where=()):
    """Итоги по строкам таблицы: число строк и суммы столбцов, всего и по группам. Для проверки расчётов."""
    cols, rows = r["columns"], r["rows"]
    if not cols:
        raise UserError("Данных в таблице нет.")
    if len(by) > 2:
        raise UserError("Группировать можно не больше чем по двум столбцам.")
    by_i = [find_column(cols, b) for b in by]
    sum_i = [find_column(cols, s) for s in sums]
    conds = parse_where(cols, where)
    # Отбор по диапазону: строки, где дата (или число) в столбце не читается, в отбор не входят; их считаем.
    ranged = {i: tests[0][1].convert for i, tests in conds.items()
              if all(hasattr(t, "convert") for _, t in tests)}
    unreadable = {i: 0 for i in ranged}
    groups, order = {}, []
    total = {"rows": 0, "sums": [0.0] * len(sum_i), "empty": [0] * len(sum_i), "bad": [[] for _ in sum_i]}
    for n, row in enumerate(rows, start=1):
        for i, convert in ranged.items():
            if convert(row[i]) is None:
                unreadable[i] += 1
        if not all(any(test(row[i]) for _, test in tests) for i, tests in conds.items()):
            continue
        # Группа без учёта регистра и лишних пробелов, как в отборе; подпись: первое написание в файле.
        shown = tuple(" ".join(str(row[i]).split()) for i in by_i)
        key = tuple(norm_text(v) for v in shown)
        g = groups.get(key)
        if g is None:
            g = groups[key] = {"rows": 0, "sums": [0.0] * len(sum_i), "label": shown}
            order.append(key)
        g["rows"] += 1
        total["rows"] += 1
        for j, i in enumerate(sum_i):
            cell = row[i]
            if cell.strip() == "":
                total["empty"][j] += 1
                continue
            v = to_number(cell)
            if v is None:
                total["bad"][j].append(n)
                continue
            g["sums"][j] += v
            total["sums"][j] += v
    return {
        "by": [cols[i] for i in by_i],
        "sums": [cols[i] for i in sum_i],
        "where": [item for tests in conds.values() for item, _ in tests],
        "data_rows": len(rows),
        "groups": [{"key": list(groups[k]["label"]), "rows": groups[k]["rows"], "sums": groups[k]["sums"]} for k in order],
        "total": total,
        "unreadable": [{"column": cols[i], "kind": "дата" if ranged[i] is to_date else "число", "rows": c}
                       for i, c in unreadable.items() if c],
    }


def fmt_number(x):
    text = f"{round(x, 6):.6f}".rstrip("0").rstrip(".")
    return "0" if text in ("-0", "") else text


def format_totals(t, r):
    out = [f"Файл: {r['file']}"]
    if r.get("type") == "xlsx":
        out.append(f"Лист: {r['sheet']['name']}")
    if t["where"]:
        out.append("Отбор: " + "; ".join(t["where"]))
    out.append(f"Строк в отборе: {t['total']['rows']} из {t['data_rows']}.")
    for u in t.get("unreadable", []):
        out.append(f"Внимание: «{u['column']}»: {u['kind']} не читается в {u['rows']} строках, в отбор не вошли.")
    if t["data_rows"] < r["data_rows"]:
        out.append(f"Внимание: в файле {r['data_rows']} строк, посчитаны первые {t['data_rows']}.")
    head = t["by"] + ["строк"] + t["sums"]
    if t["by"]:
        out.append(" | ".join(head))
        for g in t["groups"][:MAX_GROUPS]:
            out.append(" | ".join([str(k) for k in g["key"]] + [str(g["rows"])] + [fmt_number(s) for s in g["sums"]]))
        if len(t["groups"]) > MAX_GROUPS:
            out.append(f"... и ещё групп: {len(t['groups']) - MAX_GROUPS}. Сузьте отбор (--where).")
        out.append(" | ".join(["Всего"] + [""] * (len(t["by"]) - 1) + [str(t["total"]["rows"])]
                              + [fmt_number(s) for s in t["total"]["sums"]]))
    else:
        out.append("Всего строк: " + str(t["total"]["rows"]))
        for name, s in zip(t["sums"], t["total"]["sums"]):
            out.append(f"Сумма «{name}»: {fmt_number(s)}")
    for name, empty, bad in zip(t["sums"], t["total"]["empty"], t["total"]["bad"]):
        if empty:
            out.append(f"«{name}»: пустых ячеек {empty}, в сумму не вошли.")
        if bad:
            shown = ", ".join(str(n) for n in bad[:20]) + (" ..." if len(bad) > 20 else "")
            out.append(f"«{name}»: не числа в {len(bad)} строках данных ({shown}), в сумму не вошли.")
    return "\n".join(out)


def format_human(r):
    out = [f"Файл: {r['file']}"]
    if r["type"] == "csv":
        out.append(f"Кодировка: {r['encoding']}")
        out.append(f"Разделитель: {DELIM_NAMES.get(r['delimiter'], r['delimiter'])}")
    else:
        out.append("Листы:")
        for s in r["sheets"]:
            mark = " <- выбран" if s["index"] == r["sheet"]["index"] else ""
            out.append(f"  {s['index']}. {s['name']}: строк {s['rows']}{mark}")
    if r["header_row"] is None:
        out.append("Данных в таблице нет.")
        return "\n".join(out)
    levels = f" (шапка в {r['header_rows']} строки, объединённые ячейки развёрнуты)" if r["header_rows"] > 1 else ""
    out.append(f"Строка шапки: {r['header_row']}{levels}")
    out.append(f"Столбцов: {len(r['columns'])}")
    out.append("Столбцы: " + " | ".join(r["columns"]))
    out.append(f"Строк данных: {r['data_rows']}")
    for w in r.get("warnings", []):
        out.append(f"Внимание: {w}")
    out.append(f"Первые строки ({len(r['rows'])}):")
    for row in r["rows"]:
        out.append("  " + " | ".join(row))
    return "\n".join(out)


def main(argv=None):
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8")
        except Exception:
            pass
    ap = argparse.ArgumentParser(description="Заглянуть в таблицу Excel или выгрузку 1С.")
    ap.add_argument("file", help="путь к .xlsx, .csv или .txt")
    ap.add_argument("--sheet", help="номер листа (с 1) или имя; по умолчанию первый")
    ap.add_argument("--rows", type=int, default=10, help="сколько первых строк показать (10)")
    ap.add_argument("--json", action="store_true", help="вывод в JSON для скриптов")
    ap.add_argument("--sum", action="append", default=[], metavar="СТОЛБЕЦ",
                    help="итоги: сумма столбца (номер с 1, название или часть названия), можно несколько")
    ap.add_argument("--by", action="append", default=[], metavar="СТОЛБЕЦ",
                    help="итоги по группам этого столбца, до двух")
    ap.add_argument("--where", action="append", default=[], metavar="СТОЛБЕЦ=ЗНАЧЕНИЕ",
                    help="отбор строк: Столбец=значение или Столбец=от..до (даты, числа)")
    a = ap.parse_args(argv)
    totals = bool(a.sum or a.by or a.where)
    try:
        r = analyze(Path(a.file), a.sheet, LIMIT_TOTALS if totals else max(0, a.rows))
        if totals:
            t = group_totals(r, a.sum, a.by, a.where)
            print(json.dumps(t, ensure_ascii=False, indent=2) if a.json else format_totals(t, r))
            return 0
    except UserError as e:
        print(str(e), file=sys.stderr)
        if a.json:
            print(json.dumps({"error": str(e)}, ensure_ascii=False))
        return 2
    except Exception as e:  # битый файл и прочее: без трейсбека для человека
        msg = f"Не удалось прочитать файл: {type(e).__name__}: {e}"
        print(msg, file=sys.stderr)
        if a.json:
            print(json.dumps({"error": msg}, ensure_ascii=False))
        return 2
    print(json.dumps(r, ensure_ascii=False, indent=2) if a.json else format_human(r))
    return 0


if __name__ == "__main__":
    sys.exit(main())
