"""Светофор договоров: демо-процесс тренера на учебном кейсе «Дала Ойл».

Читает input/svetofor-dogovorov/reestr_dogovorov.xlsx: все листы, где есть столбцы «№ договора»
и «Дата окончания» (листы городов; «Справочник» пропускается). Считает дни до окончания договора
от даты легенды кейса 21.09.2026 и кладёт в output/svetofor-dogovorov/:
  result.json   карточка процесса на странице «Мой кабинет»;
  svetofor.xlsx тот же список для Excel.

Красный: договор уже истёк. Жёлтый: до окончания от 0 до 30 дней. Зелёный: больше 30 дней.
Нет даты окончания: «нет данных» (серый, в зелёный не попадает). Бессрочный: пометка без цвета.
Дата расчёта постоянная: это «сегодня» по легенде кейса, а не дата компьютера.

Запуск из папки кабинета: python processes/svetofor-dogovorov/run.py
"""

import json
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path

SLUG = "svetofor-dogovorov"
TITLE = "Светофор договоров"
INPUT_NAME = "reestr_dogovorov.xlsx"
CALC_DATE = date(2026, 9, 21)
YELLOW_DAYS = 30

COLUMNS = ["Город", "№ договора", "Контрагент", "Предмет договора", "Объект", "Дата окончания", "Дней до окончания"]
# Столбец результата → как он называется в реестре (сравнение без учёта регистра и лишних пробелов).
SOURCE = {
    "№ договора": "№ договора",
    "Контрагент": "контрагент",
    "Предмет договора": "предмет договора",
    "Объект": "объект",
    "Дата окончания": "дата окончания",
}
REQUIRED = ("№ договора", "Дата окончания")
NOTES = {"red": "истёк", "yellow": "до 30 дней", "green": "в порядке", "grey": "нет данных", "info": "бессрочно"}
ORDER = {"red": 0, "yellow": 1, "green": 2, "grey": 3, "info": 4}
FILLS = {"red": "FDECEB", "yellow": "FFF3D4", "green": "E8F5EB", "grey": "ECEEF1"}

NO_OPENPYXL = "Не найдена библиотека openpyxl для Excel. Скажите Claude «проверь сетап»."

ISO_DATE = re.compile(r"(\d{4})-(\d{1,2})-(\d{1,2})")
DMY_DATE = re.compile(r"(\d{1,2})\.(\d{1,2})\.(\d{4}|\d{2})(?!\d)")


class StopRun(Exception):
    """Понятная человеку причина остановки: печатается по-русски, код выхода 2 (3: файл занят Excel)."""

    def __init__(self, message, code=2):
        super().__init__(message)
        self.code = code


def find_root():
    """Корень кабинета: ближайшая папка выше этого файла, где лежит tools/read_table.py."""
    for parent in Path(__file__).resolve().parents:
        if (parent / "tools" / "read_table.py").is_file():
            return parent
    raise StopRun("Не найден tools/read_table.py: программа процесса должна лежать в папке кабинета.")


def load_reader():
    """tools/read_table.py: та же логика листов и шапок, что и у просмотра таблиц в кабинете."""
    tools = str(find_root() / "tools")
    if tools not in sys.path:
        sys.path.insert(0, tools)
    import read_table
    return read_table


def norm(text):
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def parse_end(text):
    """Дата окончания из ячейки: ('date', date) | ('open', None) | ('empty', None) | ('bad', None).

    Встречается: 2026-10-31 (дата Excel), 14.10.26, 31.12.2026, «до 09.10.2026», «бессрочно», пусто.
    """
    t = norm(text)
    if not t:
        return "empty", None
    if "бессроч" in t:
        return "open", None
    m = ISO_DATE.search(t)
    if m:
        y, mo, d = (int(g) for g in m.groups())
    else:
        m = DMY_DATE.search(t)
        if not m:
            return "bad", None
        d, mo, y = (int(g) for g in m.groups())
        if y < 100:
            y += 2000
    try:
        return "date", date(y, mo, d)
    except ValueError:
        return "bad", None


def read_contracts(path):
    """Строки всех листов, где есть «№ договора» и «Дата окончания». Город это имя листа."""
    rt = load_reader()
    path = Path(path)
    try:
        sheets = rt.analyze(path, None, 0)["sheets"]
        contracts = []
        for sheet in sheets:
            data = rt.analyze(path, str(sheet["index"]), 10 ** 6)
            cols = [norm(c) for c in data["columns"]]
            idx = {}
            for name, wanted in SOURCE.items():
                idx[name] = next((i for i, c in enumerate(cols) if wanted in c), None)
            if any(idx[name] is None for name in REQUIRED):
                continue
            for row in data["rows"]:
                if not any(row):
                    continue
                item = {name: (row[i] if i is not None else "") for name, i in idx.items()}
                item["Город"] = sheet["name"]
                contracts.append(item)
    except rt.UserError as e:
        raise StopRun(str(e))
    except ImportError:
        raise StopRun(NO_OPENPYXL)
    if not contracts:
        raise StopRun(f"В файле {path.name} нет листа со столбцами «№ договора» и «Дата окончания».")
    return contracts


def evaluate(path):
    """Договоры со статусом, отсортированные по срочности: красные, жёлтые, зелёные, без даты, бессрочные."""
    items = []
    for n, c in enumerate(read_contracts(path)):
        raw = c["Дата окончания"].strip()
        kind, end = parse_end(raw)
        days = None
        if kind == "date":
            days = (end - CALC_DATE).days
            status = "red" if days < 0 else "yellow" if days <= YELLOW_DAYS else "green"
            note = NOTES[status]
        elif kind == "open":
            status, note = "info", NOTES["info"]
        elif kind == "empty":
            status, note = "grey", NOTES["grey"]
        else:
            status, note = "grey", f"нет данных: дата «{raw}» не распознана"
        items.append({**c, "end": end, "end_text": end.strftime("%d.%m.%Y") if end else raw,
                      "days": days, "status": status, "note": note, "n": n})
    items.sort(key=lambda x: (ORDER[x["status"]], x["days"] if x["days"] is not None else 0, x["n"]))
    return items


def row_values(item):
    return [item["Город"], item["№ договора"], item["Контрагент"], item["Предмет договора"],
            item["Объект"], item["end_text"], item["days"]]


def result_from(items):
    count = {s: sum(1 for x in items if x["status"] == s) for s in ORDER}
    status = "red" if count["red"] else "yellow" if count["yellow"] else "green"
    summary = (
        f"На {CALC_DATE:%d.%m.%Y}, договоров {len(items)}. Истекли: {count['red']}. "
        f"Истекают в ближайшие {YELLOW_DAYS} дней: {count['yellow']}. В порядке: {count['green']}. "
        f"Нет данных о дате окончания: {count['grey']}. Бессрочные: {count['info']}."
    )
    return {
        "title": TITLE,
        "updated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "status": status,
        "summary": summary,
        "columns": COLUMNS,
        "rows": [row_values(x) for x in items],
        "row_status": [x["status"] for x in items],
        "row_note": [x["note"] for x in items],
        "files": [f"output/{SLUG}/svetofor.xlsx"],
    }


def build_result(path):
    return result_from(evaluate(path))


def replace_file(tmp, target):
    """Запись через временный файл: страница не увидит половину файла. Файл открыт в Excel: понятная ошибка."""
    try:
        os.replace(tmp, target)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise StopRun(f"Файл output/{SLUG}/{target.name} открыт в другой программе. "
                      "Закройте его в Excel и запустите расчёт снова.", code=3)


def write_excel(items, target):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise StopRun(NO_OPENPYXL)
    wb = Workbook()
    ws = wb.active
    ws.title = "Светофор"
    ws.append(COLUMNS + ["Статус"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for item in items:
        values = row_values(item)
        values[5] = item["end"] or item["end_text"]
        ws.append(values + [item["note"]])
        fill = FILLS.get(item["status"])
        for cell in ws[ws.max_row]:
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            if isinstance(cell.value, date):
                cell.number_format = "DD.MM.YYYY"
    for letter, width in zip("ABCDEFGH", (12, 14, 34, 24, 28, 16, 18, 14)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    tmp = target.with_name(target.name + ".tmp")
    wb.save(tmp)
    replace_file(tmp, target)


def write_json(result, target):
    tmp = target.with_name(target.name + ".tmp")
    tmp.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    replace_file(tmp, target)


def main():
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8")
        except Exception:
            pass
    try:
        root = find_root()
        src = root / "input" / SLUG / INPUT_NAME
        if not src.is_file():
            raise StopRun(f"Нет файла input/{SLUG}/{INPUT_NAME}. Положите реестр договоров в папку input/{SLUG} "
                          "(учебный файл лежит в input/case-azs/data) и запустите расчёт снова.")
        items = evaluate(src)
        result = result_from(items)
        out = root / "output" / SLUG
        out.mkdir(parents=True, exist_ok=True)
        write_excel(items, out / "svetofor.xlsx")
        write_json(result, out / "result.json")
    except StopRun as e:
        print(str(e), file=sys.stderr)
        return e.code
    print(result["summary"])
    print(f"Готово: output/{SLUG}/result.json и output/{SLUG}/svetofor.xlsx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
