"""Программа сервиса. Её запускает кнопка «Запустить» в кабинете, без Claude и без вопросов человеку.

Запуск: python main.py в папке проекта (на Mac python3 main.py).
Вход: самый свежий файл в input/. Выход: output/result.json (карточка в кабинете) и файлы результата в output/.
Понятная причина остановки: строка по-русски, код выхода 2 (3: файл результата открыт в Excel или Word).
Сбой программы: одна строка с местом сбоя, код выхода 1.

Сейчас это каркас: он показывает устройство самого свежего файла в input/. Расчёт сервиса пишется в main(),
функции каркаса остаются как есть.
"""
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent  # папка проекта: все пути только от неё
INPUT = HERE / "input"
OUT = HERE / "output"
LIMIT = 10 ** 6  # сколько строк данных читать; в файле больше: итог об этом скажет
sys.path.insert(0, str(HERE / "lib"))  # помощники из кабинета: read_table.py и llm.py


class StopRun(Exception):
    """Понятная человеку причина остановки: строка по-русски, код выхода 2 (3: файл открыт в другой программе)."""

    def __init__(self, message, code=2):
        super().__init__(message)
        self.code = code


try:
    import read_table  # таблицы: кодировка, шапка и листы, как у просмотра таблиц в кабинете
except ImportError:
    read_table = None

try:
    import llm  # модель внутри сервиса: только через него, он проверяет данные перед отправкой
except ImportError:
    llm = None


def service_name():
    """Название сервиса из паспорта service.json."""
    try:
        name = json.loads((HERE / "service.json").read_text(encoding="utf-8")).get("name")
    except (OSError, ValueError, AttributeError):
        name = None
    return name.strip() if isinstance(name, str) and name.strip() else "Сервис"


def newest_input(exts=(".csv", ".xlsx", ".txt")):
    """Самый свежий файл нужного типа в input/: новый файл за другой период подхватится сам."""
    files = [p for p in INPUT.glob("*") if p.is_file() and p.suffix.lower() in exts
             and not p.name.startswith(("~$", "."))]
    if not files:
        kinds = ", ".join(exts)
        raise StopRun(f"В папке input нет файла с данными ({kinds}). Положите файл и запустите снова.")
    # Когда файл появился в папке или изменился: копия из Проводника сохраняет старую дату изменения.
    return max(files, key=lambda p: max(p.stat().st_mtime, p.stat().st_ctime))


def read_input(src, sheet=None):
    """Таблица из файла: {"columns": [...], "rows": [[текст, ...]], "data_rows": N, "above_header": [...]}."""
    if read_table is None:
        raise StopRun("В папке lib нет файла read_table.py из кабинета. Скопируйте его из папки tools кабинета "
                      "или позовите тренера.")
    return read_table.analyze(src, sheet, LIMIT)


def number(text):
    """Число из ячейки: «1 234,5» или «1234.5». Не число (и пусто, nan, inf, 1e5): None, не ноль."""
    # split() без аргументов убирает любые пробелы, и неразрывный между разрядами в выгрузках 1С тоже.
    t = "".join(str(text if text is not None else "").split()).replace(",", ".")
    return float(t) if re.fullmatch(r"[-+]?\d+(\.\d+)?", t) else None


def column(columns, *names):
    """Номер столбца по названию из шапки (без учёта регистра и лишних пробелов). Нет: понятная остановка."""
    norm = [" ".join(str(c).split()).lower() for c in columns]
    for name in names:
        key = " ".join(name.split()).lower()
        if key in norm:
            return norm.index(key)
    shown = ", ".join(str(c) for c in columns)
    raise StopRun(f"В файле нет столбца «{names[0]}». Столбцы файла: {shown}.")


def read_env(name):
    """Ключ из файла .env в папке проекта (строка ИМЯ=значение). Значение не печатать и не писать в результат."""
    try:
        lines = (HERE / ".env").read_text(encoding="utf-8-sig").splitlines()
    except OSError:
        lines = []
    for line in lines:
        key, sep, value = line.partition("=")
        if sep and key.strip() == name:
            return value.strip().strip('"').strip("'")
    raise StopRun(f"Нет ключа {name} в файле .env в папке проекта. Добавьте строку {name}=значение и запустите снова.")


def ask_model(instruction, data=""):
    """Ответ модели или (None, причина). В модель только устройство данных и итоги, не строки файлов."""
    if llm is None:
        return None, "модель недоступна: в папке lib нет llm.py"
    try:
        return llm.ask(instruction, data), None
    except llm.LLMBlocked as e:
        return None, f"не отправлено в модель: {e}"
    except llm.LLMUnavailable as e:
        return None, f"модель недоступна: {e}"


def replace_file(tmp, target):
    """Запись через временный файл: кабинет не увидит половину файла. Файл открыт: понятная строка, код 3."""
    try:
        os.replace(tmp, target)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise StopRun(f"Файл output/{target.name} открыт в другой программе. Закройте его и запустите снова.", code=3)


def write_excel(name, sheets):
    """Файл Excel output/<name>: sheets = {"Лист": (столбцы, строки)}. Возвращает путь для files."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        raise StopRun("Не найдена библиотека openpyxl для Excel. Скажите Claude в кабинете «проверь сетап».")
    wb = Workbook()
    wb.remove(wb.active)
    for title, (columns, rows) in sheets.items():
        ws = wb.create_sheet(title[:31])
        ws.append(list(columns))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(list(row))
        for i, col in enumerate(columns, start=1):
            longest = max([len(str(col))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(longest + 2, 8), 60)
        ws.freeze_panes = "A2"
    OUT.mkdir(parents=True, exist_ok=True)
    target = OUT / name
    tmp = target.with_name(target.name + ".tmp")
    wb.save(tmp)
    replace_file(tmp, target)
    return f"output/{name}"


def write_result(result):
    """output/result.json для карточки в кабинете. NaN запрещён: он ломает карточку."""
    OUT.mkdir(parents=True, exist_ok=True)
    tmp = OUT / "result.json.tmp"
    tmp.write_text(json.dumps(result, ensure_ascii=False, allow_nan=False, indent=2), encoding="utf-8")
    replace_file(tmp, OUT / "result.json")


def now_iso():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def main():
    src = newest_input()
    table = read_input(src)
    columns, rows = table["columns"], table["rows"]  # строки данных, значения текстом

    # Каркас: устройство файла (сколько заполнено и сколько чисел в каждом столбце).
    # Здесь пишется расчёт сервиса по описанию человека; итог проверяется на известном ответе.
    card = []
    for i, name in enumerate(columns):
        values = [r[i] for r in rows if i < len(r) and str(r[i]).strip()]
        card.append([name, len(values), sum(1 for v in values if number(v) is not None)])

    summary = (f"Каркас сервиса: файл {src.name}, строк данных {len(rows)}, столбцов {len(columns)}. "
               "Опишите Claude, что должен делать сервис.")
    if table.get("data_rows", len(rows)) > len(rows):
        summary += f" Внимание: в файле {table['data_rows']} строк, прочитаны первые {len(rows)}."
    result = {
        "title": service_name(),
        "updated_at": now_iso(),
        "status": "green",
        "summary": summary,
        "columns": ["Столбец", "Заполнено", "Из них числа"],
        "rows": card,
        "files": [],
    }
    write_result(result)
    print(summary)


def where_failed(e):
    """Файл и строка программы, где случился сбой (без трейсбека: человеку одна строка)."""
    tb = e.__traceback__
    if tb is None:
        return "место неизвестно"
    while tb.tb_next is not None:
        tb = tb.tb_next
    return f"{Path(tb.tb_frame.f_code.co_filename).name}, строка {tb.tb_lineno}"


if __name__ == "__main__":
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    try:
        main()
    except StopRun as e:
        print(e, file=sys.stderr)
        sys.exit(e.code)
    except Exception as e:
        if read_table is not None and isinstance(e, read_table.UserError):
            print(e, file=sys.stderr)
            sys.exit(2)
        print(f"Сбой программы сервиса ({where_failed(e)}): {type(e).__name__}: {e}. "
              "Скажите Claude в папке проекта «почини сервис».", file=sys.stderr)
        sys.exit(1)
